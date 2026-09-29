"""
============================================================
Modulo: EXCEL / utils.py
Autor: Tiago Eneas Antunes (Z701038)
Data de Criacao: 29/09/2026
Ultima Alteracao: 29/09/2026
Versao: 1.0.0
Descricao:
    Biblioteca de funcoes utilitarias compartilhadas pelos modulos de
    automacao de arquivos Excel. O modulo centraliza conversoes de colunas,
    normalizacao de nomes, referencias, caminhos, valores e consultas de
    dimensoes utilizadas por Workbooks e Worksheets do openpyxl.
Dependencias:
    - openpyxl
Historico:
    v1.0.0 - 29/09/2026
        - Criacao inicial do modulo.
        - Inclusao de conversoes de colunas e referencias.
        - Inclusao de normalizacao de nomes de arquivos, planilhas e tabelas.
        - Inclusao de consultas de ultima linha, coluna e intervalo utilizado.
        - Inclusao de manipulacao de caminhos e nomes unicos.
============================================================
"""

# ----------------------------------------------------------------------------
# IMPORTACOES
# ----------------------------------------------------------------------------

# Importa expressoes regulares para validar referencias e normalizar nomes
import re

# Importa unicodedata para remover acentos dos textos
import unicodedata

# Importa date e datetime para normalizar valores temporais
from datetime import date, datetime

# Importa Path para trabalhar com caminhos de forma independente do sistema
from pathlib import Path

# Importa Any para anotacoes de valores flexiveis
from typing import Any

# Importa Workbook para validar arquivos Excel em memoria
from openpyxl import Workbook

# Importa MergedCell para ignorar celulas mescladas secundarias
from openpyxl.cell.cell import MergedCell

# Importa utilitarios oficiais de conversao do openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import range_boundaries

# Importa Worksheet para validar planilhas
from openpyxl.worksheet.worksheet import Worksheet


# ----------------------------------------------------------------------------
# CONSTANTES
# ----------------------------------------------------------------------------

# Define os limites estruturais do Excel
LIMITE_LINHAS_EXCEL = 1_048_576
LIMITE_COLUNAS_EXCEL = 16_384
ULTIMA_COLUNA_EXCEL = "XFD"

# Define os limites dos nomes utilizados pelo Excel
LIMITE_CARACTERES_PLANILHA = 31
LIMITE_CARACTERES_NOME_TABELA = 255

# Define as extensoes modernas suportadas pelo openpyxl
EXTENSOES_EXCEL_SUPORTADAS = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm"
}

# Define os caracteres proibidos nos nomes de planilhas e arquivos
CARACTERES_INVALIDOS_PLANILHA = {"[", "]", ":", "*", "?", "/", chr(92)}
CARACTERES_INVALIDOS_ARQUIVO = "<>:" + chr(34) + "/" + chr(92) + "|?*"

# Define o padrao utilizado para validar referencias de celulas
PADRAO_REFERENCIA_CELULA = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9]\d*)$")


# ----------------------------------------------------------------------------
# EXCECOES
# ----------------------------------------------------------------------------

class ExcelUtilsError(Exception):
    """Erro base para operacoes utilitarias do modulo Excel."""


class ValidacaoUtilsError(ExcelUtilsError, ValueError):
    """Indica que um valor nao atende as regras da funcao utilitaria."""


class CaminhoUtilsError(ExcelUtilsError, OSError):
    """Indica que um caminho de arquivo ou diretorio e invalido."""


class ReferenciaUtilsError(ExcelUtilsError, ValueError):
    """Indica que uma referencia de celula ou intervalo e invalida."""


class NomeUtilsError(ExcelUtilsError, ValueError):
    """Indica que um nome de planilha, tabela ou arquivo e invalido."""


class PlanilhaUtilsInvalidaError(ExcelUtilsError, TypeError):
    """Indica que o objeto informado nao e uma Worksheet valida."""


class WorkbookUtilsInvalidoError(ExcelUtilsError, TypeError):
    """Indica que o objeto informado nao e um Workbook valido."""


# ----------------------------------------------------------------------------
# FUNCOES INTERNAS DE VALIDACAO
# ----------------------------------------------------------------------------

def _validar_booleano(
    valor: bool,
    nome: str
) -> bool:
    """
    Valida estritamente um parametro booleano.

    Args:
        valor (bool):
            Valor que sera validado.
        nome (str):
            Nome utilizado na mensagem de erro.

    Returns:
        bool:
            O proprio valor booleano validado.

    Raises:
        TypeError:
            Caso o valor nao seja booleano.
    """
    # Impede que inteiros sejam aceitos como booleanos
    if not isinstance(valor, bool):
        raise TypeError(
            f"O parametro '{nome}' deve ser booleano."
        )

    # Retorna o valor validado
    return valor


def _validar_inteiro(
    valor: int,
    nome: str,
    minimo: int | None = None,
    maximo: int | None = None
) -> int:
    """
    Valida um numero inteiro e seus limites opcionais.

    Args:
        valor (int):
            Numero inteiro que sera validado.
        nome (str):
            Nome utilizado na mensagem de erro.
        minimo (int | None):
            Menor valor permitido.
        maximo (int | None):
            Maior valor permitido.

    Returns:
        int:
            O proprio numero inteiro validado.

    Raises:
        TypeError:
            Caso o valor nao seja inteiro.
        ValueError:
            Caso o valor esteja fora dos limites definidos.
    """
    # Rejeita booleanos porque bool herda de int em Python
    if isinstance(valor, bool) or not isinstance(valor, int):
        raise TypeError(
            f"O parametro '{nome}' deve ser inteiro."
        )

    # Verifica o limite minimo
    if minimo is not None and valor < minimo:
        raise ValueError(
            f"O parametro '{nome}' deve ser maior ou igual a {minimo}."
        )

    # Verifica o limite maximo
    if maximo is not None and valor > maximo:
        raise ValueError(
            f"O parametro '{nome}' deve ser menor ou igual a {maximo}."
        )

    # Retorna o numero validado
    return valor


def _validar_texto(
    valor: str,
    nome: str,
    permitir_vazio: bool = False
) -> str:
    """
    Valida e normaliza um parametro textual.

    Args:
        valor (str):
            Texto que sera validado.
        nome (str):
            Nome utilizado na mensagem de erro.
        permitir_vazio (bool):
            Define se uma string vazia sera permitida.

    Returns:
        str:
            Texto sem espacos externos.

    Raises:
        TypeError:
            Caso o valor nao seja uma string.
        ValueError:
            Caso o texto esteja vazio sem permissao.
    """
    # Valida a opcao booleana
    _validar_booleano(
        permitir_vazio,
        "permitir_vazio"
    )

    # Verifica o tipo do valor
    if not isinstance(valor, str):
        raise TypeError(
            f"O parametro '{nome}' deve ser uma string."
        )

    # Remove espacos das extremidades
    texto = valor.strip()

    # Rejeita texto vazio quando necessario
    if not texto and not permitir_vazio:
        raise ValueError(
            f"O parametro '{nome}' nao pode estar vazio."
        )

    # Retorna o texto normalizado
    return texto


def validar_workbook(
    workbook: Workbook
) -> Workbook:
    """
    Valida se o objeto informado e um Workbook do openpyxl.

    Args:
        workbook (Workbook):
            Workbook que sera validado.

    Returns:
        Workbook:
            O proprio Workbook validado.

    Raises:
        WorkbookUtilsInvalidoError:
            Caso o objeto nao seja um Workbook.
    """
    # Verifica o tipo exato esperado pelo modulo
    if not isinstance(workbook, Workbook):
        raise WorkbookUtilsInvalidoError(
            "O objeto fornecido nao e um Workbook do openpyxl."
        )

    # Retorna o Workbook validado
    return workbook


def validar_planilha(
    planilha: Worksheet
) -> Worksheet:
    """
    Valida se o objeto informado e uma Worksheet do openpyxl.

    Args:
        planilha (Worksheet):
            Planilha que sera validada.

    Returns:
        Worksheet:
            A propria planilha validada.

    Raises:
        PlanilhaUtilsInvalidaError:
            Caso o objeto nao seja uma Worksheet.
    """
    # Verifica o tipo da planilha
    if not isinstance(planilha, Worksheet):
        raise PlanilhaUtilsInvalidaError(
            "O objeto fornecido nao e uma planilha do openpyxl."
        )

    # Retorna a planilha validada
    return planilha


# ----------------------------------------------------------------------------
# FUNCOES DE TEXTOS E NOMES
# ----------------------------------------------------------------------------

def remover_acentos(
    texto: str
) -> str:
    """
    Remove acentos e sinais diacriticos de um texto.

    Args:
        texto (str):
            Texto que sera normalizado.

    Returns:
        str:
            Texto sem sinais diacriticos.

    Raises:
        TypeError:
            Caso texto nao seja uma string.
    """
    # Valida o texto sem impedir uma string vazia
    valor = _validar_texto(
        texto,
        "texto",
        permitir_vazio=True
    )

    # Decompoe os caracteres Unicode
    normalizado = unicodedata.normalize(
        "NFKD",
        valor
    )

    # Remove os caracteres classificados como marcas combinantes
    return "".join(
        caractere
        for caractere in normalizado
        if not unicodedata.combining(caractere)
    )


def normalizar_nome_planilha(
    nome: str,
    substituir_por: str = "_",
    truncar: bool = True
) -> str:
    """
    Normaliza um nome para uso como titulo de planilha.

    Args:
        nome (str):
            Nome original da planilha.
        substituir_por (str):
            Texto utilizado para substituir caracteres proibidos.
        truncar (bool):
            Define se nomes longos serao limitados automaticamente.

    Returns:
        str:
            Nome compativel com as regras do Excel.

    Raises:
        NomeUtilsError:
            Caso o nome final esteja vazio ou exceda o limite sem truncamento.
    """
    # Valida os parametros recebidos
    try:
        texto = _validar_texto(
            nome,
            "nome"
        )
        substituto = _validar_texto(
            substituir_por,
            "substituir_por",
            permitir_vazio=True
        )
        _validar_booleano(
            truncar,
            "truncar"
        )
    except (TypeError, ValueError) as erro:
        raise NomeUtilsError(
            "Nao foi possivel normalizar o nome da planilha."
        ) from erro

    # Substitui todos os caracteres proibidos
    for caractere in CARACTERES_INVALIDOS_PLANILHA:
        texto = texto.replace(
            caractere,
            substituto
        )

    # Remove apostrofos nas extremidades e espacos residuais
    texto = texto.strip(" '")

    # Rejeita um resultado vazio
    if not texto:
        raise NomeUtilsError(
            "O nome da planilha ficou vazio apos a normalizacao."
        )

    # Trata nomes maiores que o limite do Excel
    if len(texto) > LIMITE_CARACTERES_PLANILHA:
        if not truncar:
            raise NomeUtilsError(
                "O nome da planilha nao pode ultrapassar 31 caracteres."
            )

        texto = texto[:LIMITE_CARACTERES_PLANILHA].rstrip()

    # Retorna o nome final
    return texto


def gerar_nome_planilha_unico(
    workbook: Workbook,
    nome_base: str,
    separador: str = "_"
) -> str:
    """
    Gera um nome de planilha que ainda nao existe no Workbook.

    Args:
        workbook (Workbook):
            Workbook utilizado para verificar duplicidades.
        nome_base (str):
            Nome preferencial da nova planilha.
        separador (str):
            Separador utilizado antes do contador incremental.

    Returns:
        str:
            Nome unico e compativel com o limite do Excel.

    Raises:
        WorkbookUtilsInvalidoError:
            Caso o objeto nao seja um Workbook.
    """
    # Valida o Workbook e os textos recebidos
    validar_workbook(workbook)
    base = normalizar_nome_planilha(nome_base)
    separador_validado = _validar_texto(
        separador,
        "separador",
        permitir_vazio=True
    )

    # Retorna o nome original quando ele esta disponivel
    if base not in workbook.sheetnames:
        return base

    # Incrementa o contador ate localizar um nome disponivel
    contador = 1

    while True:
        sufixo = f"{separador_validado}{contador}"
        limite_base = LIMITE_CARACTERES_PLANILHA - len(sufixo)
        candidato = f"{base[:limite_base]}{sufixo}"

        if candidato not in workbook.sheetnames:
            return candidato

        contador += 1


def normalizar_nome_tabela(
    nome: str,
    substituir_por: str = "_"
) -> str:
    """
    Normaliza um texto para uso como nome interno de tabela.

    Args:
        nome (str):
            Nome original da tabela.
        substituir_por (str):
            Texto utilizado para substituir caracteres invalidos.

    Returns:
        str:
            Nome valido para uma tabela estruturada.

    Raises:
        NomeUtilsError:
            Caso o nome final esteja vazio ou seja reservado.
    """
    # Valida e normaliza os textos
    try:
        texto = _validar_texto(
            nome,
            "nome"
        )
        substituto = _validar_texto(
            substituir_por,
            "substituir_por",
            permitir_vazio=True
        )
    except (TypeError, ValueError) as erro:
        raise NomeUtilsError(
            "Nao foi possivel normalizar o nome da tabela."
        ) from erro

    # Remove acentos e substitui caracteres nao permitidos
    texto = remover_acentos(texto)
    texto = re.sub(
        r"[^A-Za-z0-9_.]",
        substituto,
        texto
    )

    # Reduz repeticoes do substituto quando ele nao esta vazio
    if substituto:
        texto = re.sub(
            re.escape(substituto) + r"+",
            substituto,
            texto
        )

    # Remove separadores nas extremidades
    texto = texto.strip("_.")

    # Adiciona prefixo quando o primeiro caractere e numerico
    if texto and texto[0].isdigit():
        texto = f"Tabela_{texto}"

    # Rejeita nomes vazios ou reservados
    if not texto or texto.upper() in {"R", "C"}:
        raise NomeUtilsError(
            "O nome final da tabela e invalido ou reservado."
        )

    # Aplica o limite preventivo definido pelo modulo
    return texto[:LIMITE_CARACTERES_NOME_TABELA]


def sanitizar_nome_arquivo(
    nome: str,
    substituto: str = "_",
    tamanho_maximo: int = 255
) -> str:
    """
    Remove caracteres invalidos de um nome de arquivo.

    Args:
        nome (str):
            Nome original do arquivo.
        substituto (str):
            Texto utilizado no lugar dos caracteres proibidos.
        tamanho_maximo (int):
            Tamanho total maximo do nome, incluindo a extensao.

    Returns:
        str:
            Nome de arquivo sanitizado.

    Raises:
        NomeUtilsError:
            Caso o nome final esteja vazio ou o limite seja insuficiente.
    """
    # Valida os parametros basicos
    try:
        texto = _validar_texto(
            nome,
            "nome"
        )
        substituto_validado = _validar_texto(
            substituto,
            "substituto",
            permitir_vazio=True
        )
        limite = _validar_inteiro(
            tamanho_maximo,
            "tamanho_maximo",
            minimo=1
        )
    except (TypeError, ValueError) as erro:
        raise NomeUtilsError(
            "Nao foi possivel sanitizar o nome do arquivo."
        ) from erro

    # Substitui os caracteres proibidos no Windows
    traducao = str.maketrans({
        caractere: substituto_validado
        for caractere in CARACTERES_INVALIDOS_ARQUIVO
    })
    texto = texto.translate(traducao).rstrip(" .")

    # Separa nome e extensao para preservar o sufixo durante o corte
    caminho = Path(texto)
    extensao = caminho.suffix
    base = caminho.stem.strip(" ._")

    if not base:
        base = "arquivo"

    if len(extensao) >= limite:
        raise NomeUtilsError(
            "O tamanho maximo e insuficiente para preservar a extensao."
        )

    # Reduz somente a parte principal do nome
    base = base[: limite - len(extensao)].rstrip(" ._") or "arquivo"
    resultado = f"{base}{extensao}"

    # Garante que o resultado respeite o limite
    return resultado[:limite]


# ----------------------------------------------------------------------------
# FUNCOES DE COLUNAS E REFERENCIAS
# ----------------------------------------------------------------------------

def converter_coluna_para_numero(
    coluna: str
) -> int:
    """
    Converte uma letra de coluna para seu indice numerico.

    Args:
        coluna (str):
            Letra da coluna entre A e XFD.

    Returns:
        int:
            Indice da coluna iniciado em 1.

    Raises:
        ValidacaoUtilsError:
            Caso a coluna esteja fora dos limites do Excel.
    """
    # Valida e normaliza a letra da coluna
    try:
        letra = _validar_texto(
            coluna,
            "coluna"
        ).upper().replace("$", "")
        numero = column_index_from_string(letra)
    except (TypeError, ValueError) as erro:
        raise ValidacaoUtilsError(
            f"A coluna informada e invalida: '{coluna}'."
        ) from erro

    # Verifica o limite estrutural do Excel
    if numero > LIMITE_COLUNAS_EXCEL:
        raise ValidacaoUtilsError(
            f"A coluna nao pode ultrapassar {ULTIMA_COLUNA_EXCEL}."
        )

    # Retorna o indice calculado
    return numero


def converter_numero_para_coluna(
    numero: int
) -> str:
    """
    Converte um indice numerico para a letra da coluna correspondente.

    Args:
        numero (int):
            Indice da coluna iniciado em 1.

    Returns:
        str:
            Letra da coluna entre A e XFD.

    Raises:
        ValueError:
            Caso o numero esteja fora dos limites do Excel.
    """
    # Valida o indice dentro dos limites da aplicacao
    indice = _validar_inteiro(
        numero,
        "numero",
        minimo=1,
        maximo=LIMITE_COLUNAS_EXCEL
    )

    # Converte o indice utilizando o utilitario oficial
    return get_column_letter(indice)


def normalizar_referencia_celula(
    referencia: str,
    remover_absolutos: bool = True
) -> str:
    """
    Valida e normaliza uma referencia de celula do Excel.

    Args:
        referencia (str):
            Referencia como A1, $A$1, A$1 ou $A1.
        remover_absolutos (bool):
            Define se os sinais de dolar serao removidos.

    Returns:
        str:
            Referencia validada em letras maiusculas.

    Raises:
        ReferenciaUtilsError:
            Caso a referencia esteja fora dos limites do Excel.
    """
    # Valida os parametros recebidos
    try:
        texto = _validar_texto(
            referencia,
            "referencia"
        ).upper()
        _validar_booleano(
            remover_absolutos,
            "remover_absolutos"
        )
    except (TypeError, ValueError) as erro:
        raise ReferenciaUtilsError(
            "A referencia da celula e invalida."
        ) from erro

    # Extrai a coluna e a linha da referencia
    correspondencia = PADRAO_REFERENCIA_CELULA.fullmatch(texto)

    if correspondencia is None:
        raise ReferenciaUtilsError(
            f"A referencia da celula e invalida: '{referencia}'."
        )

    coluna, linha_texto = correspondencia.groups()
    numero_coluna = converter_coluna_para_numero(coluna)
    numero_linha = int(linha_texto)

    if numero_linha > LIMITE_LINHAS_EXCEL:
        raise ReferenciaUtilsError(
            f"A linha nao pode ultrapassar {LIMITE_LINHAS_EXCEL}."
        )

    # Retorna a forma simples quando solicitado
    if remover_absolutos:
        return f"{converter_numero_para_coluna(numero_coluna)}{numero_linha}"

    # Preserva os marcadores absolutos originais
    return texto


def normalizar_referencia_intervalo(
    referencia: str
) -> str:
    """
    Valida e normaliza uma referencia de intervalo do Excel.

    Args:
        referencia (str):
            Referencia de uma celula ou intervalo, como A1:D10.

    Returns:
        str:
            Intervalo normalizado em letras maiusculas.

    Raises:
        ReferenciaUtilsError:
            Caso os limites sejam invalidos ou estejam invertidos.
    """
    # Valida o texto da referencia
    try:
        texto = _validar_texto(
            referencia,
            "referencia"
        ).upper().replace("$", "")
        coluna_inicial, linha_inicial, coluna_final, linha_final = (
            range_boundaries(texto)
        )
    except (TypeError, ValueError) as erro:
        raise ReferenciaUtilsError(
            f"O intervalo informado e invalido: '{referencia}'."
        ) from erro

    # Verifica limites ausentes, zero e inversoes
    limites = (
        coluna_inicial,
        linha_inicial,
        coluna_final,
        linha_final
    )

    if any(limite is None for limite in limites) or min(limites) < 1:
        raise ReferenciaUtilsError(
            f"O intervalo informado e invalido: '{referencia}'."
        )

    if coluna_inicial > coluna_final or linha_inicial > linha_final:
        raise ReferenciaUtilsError(
            "Os limites inicial e final do intervalo estao invertidos."
        )

    if coluna_final > LIMITE_COLUNAS_EXCEL or linha_final > LIMITE_LINHAS_EXCEL:
        raise ReferenciaUtilsError(
            "O intervalo ultrapassa os limites estruturais do Excel."
        )

    # Monta a forma canonica do intervalo
    inicio = f"{get_column_letter(coluna_inicial)}{linha_inicial}"
    fim = f"{get_column_letter(coluna_final)}{linha_final}"

    return inicio if inicio == fim else f"{inicio}:{fim}"


def obter_limites_intervalo(
    referencia: str
) -> dict[str, int | str]:
    """
    Retorna os limites e dimensoes de um intervalo.

    Args:
        referencia (str):
            Referencia de uma celula ou intervalo do Excel.

    Returns:
        dict[str, int | str]:
            Limites, quantidades e referencia normalizada.

    Raises:
        ReferenciaUtilsError:
            Caso o intervalo seja invalido.
    """
    # Normaliza a referencia antes de obter os limites
    intervalo = normalizar_referencia_intervalo(referencia)
    coluna_inicial, linha_inicial, coluna_final, linha_final = (
        range_boundaries(intervalo)
    )

    # Retorna os limites estruturados
    return {
        "referencia": intervalo,
        "coluna_inicial": coluna_inicial,
        "linha_inicial": linha_inicial,
        "coluna_final": coluna_final,
        "linha_final": linha_final,
        "quantidade_colunas": coluna_final - coluna_inicial + 1,
        "quantidade_linhas": linha_final - linha_inicial + 1,
        "quantidade_celulas": (
            (coluna_final - coluna_inicial + 1)
            * (linha_final - linha_inicial + 1)
        )
    }


# ----------------------------------------------------------------------------
# FUNCOES DE DIMENSOES E CELULAS
# ----------------------------------------------------------------------------

def celula_esta_vazia(
    valor: Any,
    considerar_espacos: bool = True
) -> bool:
    """
    Verifica se um valor deve ser tratado como vazio.

    Args:
        valor (Any):
            Valor que sera avaliado.
        considerar_espacos (bool):
            Define se strings compostas por espacos serao vazias.

    Returns:
        bool:
            True quando o valor deve ser considerado vazio.

    Raises:
        TypeError:
            Caso considerar_espacos nao seja booleano.
    """
    # Valida a configuracao
    _validar_booleano(
        considerar_espacos,
        "considerar_espacos"
    )

    # None sempre representa uma celula vazia
    if valor is None:
        return True

    # Trata strings vazias ou compostas por espacos
    if isinstance(valor, str):
        return not valor.strip() if considerar_espacos else valor == ""

    # Outros valores sao considerados preenchidos
    return False


def obter_ultima_linha(
    planilha: Worksheet,
    coluna: str | int | None = None,
    ignorar_espacos: bool = True
) -> int:
    """
    Localiza a ultima linha preenchida de uma planilha ou coluna.

    Args:
        planilha (Worksheet):
            Planilha que sera consultada.
        coluna (str | int | None):
            Coluna especifica ou None para analisar toda a planilha.
        ignorar_espacos (bool):
            Define se textos somente com espacos serao vazios.

    Returns:
        int:
            Ultima linha preenchida ou zero quando nao houver dados.

    Raises:
        PlanilhaUtilsInvalidaError:
            Caso o objeto nao seja uma Worksheet.
    """
    # Valida a planilha e a configuracao
    validar_planilha(planilha)
    _validar_booleano(
        ignorar_espacos,
        "ignorar_espacos"
    )

    # Determina as colunas que serao verificadas
    if coluna is None:
        colunas = range(1, planilha.max_column + 1)
    elif isinstance(coluna, str):
        colunas = [converter_coluna_para_numero(coluna)]
    else:
        colunas = [
            _validar_inteiro(
                coluna,
                "coluna",
                minimo=1,
                maximo=LIMITE_COLUNAS_EXCEL
            )
        ]

    # Percorre as linhas de baixo para cima
    for linha in range(planilha.max_row, 0, -1):
        for indice_coluna in colunas:
            valor = planilha.cell(
                row=linha,
                column=indice_coluna
            ).value

            if not celula_esta_vazia(
                valor,
                considerar_espacos=ignorar_espacos
            ):
                return linha

    # Retorna zero quando nao existem dados
    return 0


def obter_ultima_coluna(
    planilha: Worksheet,
    linha: int | None = None,
    ignorar_espacos: bool = True
) -> int:
    """
    Localiza a ultima coluna preenchida de uma planilha ou linha.

    Args:
        planilha (Worksheet):
            Planilha que sera consultada.
        linha (int | None):
            Linha especifica ou None para analisar toda a planilha.
        ignorar_espacos (bool):
            Define se textos somente com espacos serao vazios.

    Returns:
        int:
            Ultima coluna preenchida ou zero quando nao houver dados.

    Raises:
        ValueError:
            Caso a linha esteja fora dos limites do Excel.
    """
    # Valida a planilha e a configuracao
    validar_planilha(planilha)
    _validar_booleano(
        ignorar_espacos,
        "ignorar_espacos"
    )

    # Determina as linhas que serao analisadas
    if linha is None:
        linhas = range(1, planilha.max_row + 1)
    else:
        linhas = [
            _validar_inteiro(
                linha,
                "linha",
                minimo=1,
                maximo=LIMITE_LINHAS_EXCEL
            )
        ]

    # Percorre as colunas da direita para a esquerda
    for coluna in range(planilha.max_column, 0, -1):
        for indice_linha in linhas:
            valor = planilha.cell(
                row=indice_linha,
                column=coluna
            ).value

            if not celula_esta_vazia(
                valor,
                considerar_espacos=ignorar_espacos
            ):
                return coluna

    # Retorna zero quando nao existem dados
    return 0


def obter_intervalo_utilizado(
    planilha: Worksheet,
    ignorar_espacos: bool = True
) -> str | None:
    """
    Calcula o menor intervalo retangular que contem os dados da planilha.

    Args:
        planilha (Worksheet):
            Planilha que sera consultada.
        ignorar_espacos (bool):
            Define se textos somente com espacos serao vazios.

    Returns:
        str | None:
            Intervalo utilizado ou None quando a planilha estiver vazia.

    Raises:
        PlanilhaUtilsInvalidaError:
            Caso o objeto nao seja uma Worksheet.
    """
    # Valida a planilha e a configuracao
    validar_planilha(planilha)
    _validar_booleano(
        ignorar_espacos,
        "ignorar_espacos"
    )

    # Inicializa os limites encontrados
    linhas: list[int] = []
    colunas: list[int] = []

    # Percorre somente as celulas materializadas na planilha
    for linha in planilha.iter_rows():
        for celula in linha:
            if isinstance(celula, MergedCell):
                continue

            if not celula_esta_vazia(
                celula.value,
                considerar_espacos=ignorar_espacos
            ):
                linhas.append(celula.row)
                colunas.append(celula.column)

    # Retorna None quando nao existem valores
    if not linhas:
        return None

    # Monta o menor retangulo que contem todas as celulas preenchidas
    inicio = f"{get_column_letter(min(colunas))}{min(linhas)}"
    fim = f"{get_column_letter(max(colunas))}{max(linhas)}"

    return inicio if inicio == fim else f"{inicio}:{fim}"


# ----------------------------------------------------------------------------
# FUNCOES DE CAMINHOS
# ----------------------------------------------------------------------------

def normalizar_caminho(
    caminho: str | Path
) -> Path:
    """
    Converte uma string ou Path em caminho absoluto normalizado.

    Args:
        caminho (str | Path):
            Caminho que sera normalizado.

    Returns:
        Path:
            Caminho absoluto com o usuario expandido.

    Raises:
        CaminhoUtilsError:
            Caso o caminho possua tipo invalido ou esteja vazio.
    """
    # Valida o tipo do caminho
    if not isinstance(caminho, (str, Path)):
        raise CaminhoUtilsError(
            "O caminho deve ser uma string ou objeto Path."
        )

    # Rejeita uma string vazia
    if isinstance(caminho, str) and not caminho.strip():
        raise CaminhoUtilsError(
            "O caminho nao pode estar vazio."
        )

    # Expande o usuario e resolve a forma absoluta sem exigir existencia
    try:
        return Path(caminho).expanduser().resolve()
    except (OSError, RuntimeError) as erro:
        raise CaminhoUtilsError(
            f"Nao foi possivel normalizar o caminho: '{caminho}'."
        ) from erro


def validar_extensao_excel(
    caminho: str | Path,
    extensoes: set[str] | None = None
) -> Path:
    """
    Valida se um caminho possui uma extensao Excel permitida.

    Args:
        caminho (str | Path):
            Caminho do arquivo que sera validado.
        extensoes (set[str] | None):
            Extensoes permitidas ou None para usar o conjunto padrao.

    Returns:
        Path:
            Caminho absoluto validado.

    Raises:
        CaminhoUtilsError:
            Caso a extensao nao seja permitida.
    """
    # Normaliza o caminho recebido
    arquivo = normalizar_caminho(caminho)

    # Define e normaliza as extensoes permitidas
    permitidas = EXTENSOES_EXCEL_SUPORTADAS if extensoes is None else extensoes

    if not isinstance(permitidas, set) or not permitidas:
        raise CaminhoUtilsError(
            "O parametro 'extensoes' deve ser um conjunto nao vazio ou None."
        )

    normalizadas = {
        extensao.lower()
        if extensao.startswith(".")
        else f".{extensao.lower()}"
        for extensao in permitidas
    }

    # Verifica a extensao atual do caminho
    if arquivo.suffix.lower() not in normalizadas:
        raise CaminhoUtilsError(
            f"A extensao '{arquivo.suffix}' nao e permitida. "
            f"Extensoes aceitas: {sorted(normalizadas)}."
        )

    # Retorna o caminho validado
    return arquivo


def garantir_diretorio(
    diretorio: str | Path
) -> Path:
    """
    Cria um diretorio e seus pais quando eles ainda nao existem.

    Args:
        diretorio (str | Path):
            Caminho do diretorio que sera garantido.

    Returns:
        Path:
            Diretorio absoluto existente.

    Raises:
        CaminhoUtilsError:
            Caso o caminho existente seja um arquivo ou a criacao falhe.
    """
    # Normaliza o caminho do diretorio
    caminho = normalizar_caminho(diretorio)

    # Rejeita um arquivo existente usado como diretorio
    if caminho.exists() and not caminho.is_dir():
        raise CaminhoUtilsError(
            f"O caminho informado nao e um diretorio: '{caminho}'."
        )

    # Cria toda a hierarquia necessaria
    try:
        caminho.mkdir(
            parents=True,
            exist_ok=True
        )
    except OSError as erro:
        raise CaminhoUtilsError(
            f"Nao foi possivel criar o diretorio: '{caminho}'."
        ) from erro

    # Retorna o diretorio pronto para uso
    return caminho


def gerar_caminho_unico(
    caminho: str | Path,
    separador: str = "_"
) -> Path:
    """
    Gera um caminho que nao conflita com arquivo ou diretorio existente.

    Args:
        caminho (str | Path):
            Caminho preferencial do arquivo ou diretorio.
        separador (str):
            Separador utilizado antes do contador incremental.

    Returns:
        Path:
            Caminho original ou uma versao numerada disponivel.

    Raises:
        CaminhoUtilsError:
            Caso o caminho ou separador sejam invalidos.
    """
    # Normaliza o caminho e o separador
    destino = normalizar_caminho(caminho)
    try:
        separador_validado = _validar_texto(
            separador,
            "separador",
            permitir_vazio=True
        )
    except (TypeError, ValueError) as erro:
        raise CaminhoUtilsError(
            "O separador utilizado no caminho e invalido."
        ) from erro

    # Retorna o caminho original quando ele esta disponivel
    if not destino.exists():
        return destino

    # Preserva o sufixo do arquivo durante a numeracao
    contador = 1

    while True:
        candidato = destino.with_name(
            f"{destino.stem}{separador_validado}{contador}{destino.suffix}"
        )

        if not candidato.exists():
            return candidato

        contador += 1


# ----------------------------------------------------------------------------
# FUNCOES DE VALORES
# ----------------------------------------------------------------------------

def normalizar_valor_excel(
    valor: Any,
    vazio_como_none: bool = True
) -> Any:
    """
    Normaliza valores antes de escreve-los em uma celula Excel.

    Args:
        valor (Any):
            Valor original que sera normalizado.
        vazio_como_none (bool):
            Define se strings vazias serao convertidas para None.

    Returns:
        Any:
            Valor compativel com a escrita pelo openpyxl.

    Raises:
        TypeError:
            Caso vazio_como_none nao seja booleano.
    """
    # Valida a configuracao booleana
    _validar_booleano(
        vazio_como_none,
        "vazio_como_none"
    )

    # Converte objetos Path em texto
    if isinstance(valor, Path):
        return str(valor)

    # Preserva datas e horarios suportados pelo openpyxl
    if isinstance(valor, (date, datetime)):
        return valor

    # Converte colecoes simples para texto legivel
    if isinstance(valor, (list, tuple, set)):
        return "; ".join(
            str(item)
            for item in valor
        )

    # Converte dicionarios para sua representacao textual
    if isinstance(valor, dict):
        return str(valor)

    # Trata strings vazias conforme a configuracao
    if isinstance(valor, str):
        texto = valor.strip()

        if not texto and vazio_como_none:
            return None

        return texto

    # Retorna tipos escalares sem alteracao
    return valor


# ----------------------------------------------------------------------------
# EXPORTACOES PUBLICAS
# ----------------------------------------------------------------------------

__all__ = [
    "LIMITE_LINHAS_EXCEL",
    "LIMITE_COLUNAS_EXCEL",
    "ULTIMA_COLUNA_EXCEL",
    "LIMITE_CARACTERES_PLANILHA",
    "LIMITE_CARACTERES_NOME_TABELA",
    "EXTENSOES_EXCEL_SUPORTADAS",
    "CARACTERES_INVALIDOS_PLANILHA",
    "CARACTERES_INVALIDOS_ARQUIVO",
    "ExcelUtilsError",
    "ValidacaoUtilsError",
    "CaminhoUtilsError",
    "ReferenciaUtilsError",
    "NomeUtilsError",
    "PlanilhaUtilsInvalidaError",
    "WorkbookUtilsInvalidoError",
    "validar_workbook",
    "validar_planilha",
    "remover_acentos",
    "normalizar_nome_planilha",
    "gerar_nome_planilha_unico",
    "normalizar_nome_tabela",
    "sanitizar_nome_arquivo",
    "converter_coluna_para_numero",
    "converter_numero_para_coluna",
    "normalizar_referencia_celula",
    "normalizar_referencia_intervalo",
    "obter_limites_intervalo",
    "celula_esta_vazia",
    "obter_ultima_linha",
    "obter_ultima_coluna",
    "obter_intervalo_utilizado",
    "normalizar_caminho",
    "validar_extensao_excel",
    "garantir_diretorio",
    "gerar_caminho_unico",
    "normalizar_valor_excel"
]
