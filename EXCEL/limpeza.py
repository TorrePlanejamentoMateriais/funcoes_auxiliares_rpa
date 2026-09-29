"""
============================================================
Módulo: EXCEL / limpeza.py
Autor: Tiago Eneas Antunes (Z701038)
Data de Criação: 29/09/2026
Última Alteração: 29/09/2026
Versão: 1.0.0
Descrição:
    Biblioteca de funções auxiliares para limpeza de dados em DataFrames
    e planilhas Excel. O módulo permite remover espaços, normalizar
    textos e cabeçalhos, substituir valores, tratar dados ausentes,
    remover duplicidades, excluir linhas e colunas vazias e limpar
    células ou intervalos utilizando pandas e openpyxl.
Funções disponíveis:
    - limpar_texto()
    - normalizar_cabecalho()
    - limpar_espacos_dataframe()
    - normalizar_cabecalhos_dataframe()
    - substituir_valores_dataframe()
    - preencher_valores_ausentes()
    - remover_linhas_vazias_dataframe()
    - remover_colunas_vazias_dataframe()
    - remover_duplicidades_dataframe()
    - remover_linhas_por_valor()
    - limpar_celula()
    - limpar_intervalo()
    - limpar_espacos_planilha()
    - substituir_valores_planilha()
    - remover_linhas_vazias_planilha()
    - remover_colunas_vazias_planilha()
    - remover_linhas_duplicadas_planilha()
    - limpar_planilha()
Dependências:
    - pandas
    - openpyxl
Histórico:
    v1.0.0 - 29/09/2026
        - Criação inicial do módulo.
        - Inclusão da limpeza de textos e cabeçalhos.
        - Inclusão das operações de limpeza para DataFrames.
        - Inclusão das operações de limpeza para planilhas Excel.
        - Inclusão da remoção de vazios e duplicidades.
============================================================
"""

# ----------------------------------------------------------------------------
# IMPORTAÇÕES
# ----------------------------------------------------------------------------

# Importa expressões regulares para normalização de textos
import re

# Importa unicodedata para remoção opcional de acentos
import unicodedata

# Importa Any, Iterable e Mapping para anotações flexíveis
from collections.abc import Iterable, Mapping
from typing import Any

# Importa pandas para limpeza e transformação de DataFrames
import pandas as pd

# Importa MergedCell para identificar células mescladas secundárias
from openpyxl.cell.cell import Cell, MergedCell

# Importa range_boundaries para interpretar intervalos do Excel
from openpyxl.utils import range_boundaries

# Importa Worksheet para validar objetos de planilha
from openpyxl.worksheet.worksheet import Worksheet


# ----------------------------------------------------------------------------
# CONSTANTES
# ----------------------------------------------------------------------------

# Define o padrão utilizado para localizar sequências de espaços
PADRAO_ESPACOS = re.compile(r"\s+")

# Define o caractere padrão usado na normalização de cabeçalhos
SEPARADOR_CABECALHO_PADRAO = "_"

# Define os valores textuais tratados como ausentes quando solicitado
TEXTOS_AUSENTES_PADRAO = {
    "",
    "none",
    "null",
    "nan",
    "nat",
    "n/a",
    "na",
    "-"
}


# ----------------------------------------------------------------------------
# EXCEÇÕES
# ----------------------------------------------------------------------------

class ExcelLimpezaError(Exception):
    """Erro base para operações de limpeza de dados Excel."""


class DataFrameLimpezaInvalidoError(ExcelLimpezaError, TypeError):
    """Indica que o objeto informado não é um DataFrame válido."""


class PlanilhaLimpezaInvalidaError(ExcelLimpezaError, TypeError):
    """Indica que o objeto informado não é uma Worksheet válida."""


class ReferenciaLimpezaInvalidaError(ExcelLimpezaError, ValueError):
    """Indica que uma referência de célula ou intervalo é inválida."""


class ColunaLimpezaNaoEncontradaError(ExcelLimpezaError, KeyError):
    """Indica que uma coluna solicitada não existe no DataFrame."""


class ConfiguracaoLimpezaInvalidaError(ExcelLimpezaError, ValueError):
    """Indica que uma configuração de limpeza possui valor inválido."""


# ----------------------------------------------------------------------------
# FUNÇÕES INTERNAS DE VALIDAÇÃO
# ----------------------------------------------------------------------------

def _validar_booleano(
    valor: bool,
    nome: str
) -> bool:
    """
    Valida estritamente um parâmetro booleano.

    Args:
        valor (bool):
            Valor que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.

    Returns:
        bool:
            O próprio valor validado.

    Raises:
        TypeError:
            Caso o valor não seja booleano.
    """
    # Verifica se o valor informado é booleano
    if not isinstance(valor, bool):
        raise TypeError(
            f"O parâmetro '{nome}' deve ser booleano."
        )

    # Retorna o valor validado
    return valor


def _validar_texto(
    valor: str,
    nome: str,
    permitir_vazio: bool = False
) -> str:
    """
    Valida e normaliza um parâmetro textual.

    Args:
        valor (str):
            Texto que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        permitir_vazio (bool):
            Define se uma string vazia será aceita.

    Returns:
        str:
            Texto sem espaços externos.

    Raises:
        TypeError:
            Caso o valor não seja uma string.
        ValueError:
            Caso o texto esteja vazio e isso não seja permitido.
    """
    # Valida a opção de texto vazio
    _validar_booleano(
        permitir_vazio,
        "permitir_vazio"
    )

    # Verifica se o valor informado é uma string
    if not isinstance(valor, str):
        raise TypeError(
            f"O parâmetro '{nome}' deve ser uma string."
        )

    # Remove os espaços externos
    texto = valor.strip()

    # Verifica se o conteúdo vazio é permitido
    if not texto and not permitir_vazio:
        raise ValueError(
            f"O parâmetro '{nome}' não pode estar vazio."
        )

    # Retorna o texto normalizado
    return texto


def _validar_dataframe(
    dataframe: pd.DataFrame
) -> pd.DataFrame:
    """
    Valida se o objeto informado é um DataFrame do pandas.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será validado.

    Returns:
        pd.DataFrame:
            O próprio DataFrame validado.

    Raises:
        DataFrameLimpezaInvalidoError:
            Caso o objeto não seja um DataFrame.
    """
    # Verifica o tipo do objeto recebido
    if not isinstance(dataframe, pd.DataFrame):
        raise DataFrameLimpezaInvalidoError(
            "O objeto fornecido não é um DataFrame do pandas."
        )

    # Retorna o DataFrame validado
    return dataframe


def _validar_planilha(
    planilha: Worksheet
) -> Worksheet:
    """
    Valida se o objeto informado é uma Worksheet do openpyxl.

    Args:
        planilha (Worksheet):
            Planilha que será validada.

    Returns:
        Worksheet:
            A própria planilha validada.

    Raises:
        PlanilhaLimpezaInvalidaError:
            Caso o objeto não seja uma Worksheet.
    """
    # Verifica o tipo do objeto recebido
    if not isinstance(planilha, Worksheet):
        raise PlanilhaLimpezaInvalidaError(
            "O objeto fornecido não é uma planilha do openpyxl."
        )

    # Retorna a planilha validada
    return planilha


def _normalizar_intervalo(
    referencia: str
) -> tuple[str, tuple[int, int, int, int]]:
    """
    Valida uma referência e retorna seus limites numéricos.

    Args:
        referencia (str):
            Referência de célula ou intervalo no padrão do Excel.

    Returns:
        tuple[str, tuple[int, int, int, int]]:
            Referência normalizada e limites numéricos.

    Raises:
        ReferenciaLimpezaInvalidaError:
            Caso a referência seja incompleta ou inválida.
    """
    # Valida e normaliza o texto da referência
    intervalo = _validar_texto(
        referencia,
        "referencia"
    ).upper()

    # Tenta obter os limites da referência
    try:
        limites = range_boundaries(intervalo)
    except (TypeError, ValueError) as erro:
        raise ReferenciaLimpezaInvalidaError(
            f"A referência informada é inválida: '{referencia}'."
        ) from erro

    # Rejeita referências parciais ou com índice menor que um
    if any(limite is None for limite in limites) or min(limites) < 1:
        raise ReferenciaLimpezaInvalidaError(
            f"A referência informada é inválida: '{referencia}'."
        )

    # Retorna a referência e seus limites
    return intervalo, limites


def _validar_colunas_dataframe(
    dataframe: pd.DataFrame,
    colunas: Iterable[str] | None
) -> list[str]:
    """
    Valida uma seleção opcional de colunas do DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que contém as colunas.
        colunas (Iterable[str] | None):
            Coleção opcional de colunas selecionadas.

    Returns:
        list[str]:
            Colunas validadas na ordem informada.

    Raises:
        TypeError:
            Caso colunas não seja um iterável não textual.
        ColunaLimpezaNaoEncontradaError:
            Caso alguma coluna não exista no DataFrame.
    """
    # Retorna todas as colunas quando nenhuma seleção foi informada
    if colunas is None:
        return list(dataframe.columns)

    # Rejeita uma string isolada
    if isinstance(colunas, (str, bytes)) or not isinstance(
        colunas,
        Iterable
    ):
        raise TypeError(
            "O parâmetro 'colunas' deve ser um iterável não textual ou None."
        )

    # Converte a coleção para lista
    resultado = list(colunas)

    # Verifica se todos os nomes são strings
    if not all(isinstance(coluna, str) for coluna in resultado):
        raise TypeError(
            "Todos os itens de 'colunas' devem ser strings."
        )

    # Identifica as colunas ausentes
    ausentes = [
        coluna
        for coluna in resultado
        if coluna not in dataframe.columns
    ]

    # Rejeita colunas inexistentes
    if ausentes:
        raise ColunaLimpezaNaoEncontradaError(
            f"As colunas não foram encontradas: {ausentes}."
        )

    # Retorna as colunas validadas
    return resultado


# ----------------------------------------------------------------------------
# FUNÇÕES DE LIMPEZA DE TEXTOS
# ----------------------------------------------------------------------------

def limpar_texto(
    valor: Any,
    remover_espacos_externos: bool = True,
    reduzir_espacos: bool = True,
    remover_quebras: bool = True,
    remover_acentos: bool = False,
    converter_minusculas: bool = False,
    converter_maiusculas: bool = False,
    vazio_como_none: bool = False
) -> Any:
    """
    Limpa e normaliza um valor textual sem alterar valores não textuais.

    Args:
        valor (Any):
            Valor que será limpo.
        remover_espacos_externos (bool):
            Define se os espaços das extremidades serão removidos.
        reduzir_espacos (bool):
            Define se sequências de espaços serão reduzidas para um espaço.
        remover_quebras (bool):
            Define se tabulações e quebras de linha serão substituídas.
        remover_acentos (bool):
            Define se os acentos serão removidos.
        converter_minusculas (bool):
            Define se o texto será convertido para letras minúsculas.
        converter_maiusculas (bool):
            Define se o texto será convertido para letras maiúsculas.
        vazio_como_none (bool):
            Define se textos vazios serão convertidos para None.

    Returns:
        Any:
            Texto limpo ou o valor original quando ele não for textual.

    Raises:
        ConfiguracaoLimpezaInvalidaError:
            Caso as conversões para maiúsculas e minúsculas sejam simultâneas.
    """
    # Valida as opções booleanas
    for nome, opcao in {
        "remover_espacos_externos": remover_espacos_externos,
        "reduzir_espacos": reduzir_espacos,
        "remover_quebras": remover_quebras,
        "remover_acentos": remover_acentos,
        "converter_minusculas": converter_minusculas,
        "converter_maiusculas": converter_maiusculas,
        "vazio_como_none": vazio_como_none
    }.items():
        _validar_booleano(
            opcao,
            nome
        )

    # Impede configurações contraditórias
    if converter_minusculas and converter_maiusculas:
        raise ConfiguracaoLimpezaInvalidaError(
            "Não é possível converter o texto para maiúsculas e minúsculas "
            "na mesma operação."
        )

    # Mantém valores que não são strings
    if not isinstance(valor, str):
        return valor

    # Copia o texto recebido
    texto = valor

    # Substitui caracteres de quebra por espaços
    if remover_quebras:
        texto = texto.replace("\r", " ").replace("\n", " ").replace("\t", " ")

    # Reduz sequências de espaços
    if reduzir_espacos:
        texto = PADRAO_ESPACOS.sub(" ", texto)

    # Remove espaços das extremidades
    if remover_espacos_externos:
        texto = texto.strip()

    # Remove acentos quando solicitado
    if remover_acentos:
        texto = "".join(
            caractere
            for caractere in unicodedata.normalize("NFKD", texto)
            if not unicodedata.combining(caractere)
        )

    # Converte o texto para a caixa solicitada
    if converter_minusculas:
        texto = texto.lower()
    elif converter_maiusculas:
        texto = texto.upper()

    # Converte texto vazio para None quando solicitado
    if vazio_como_none and not texto:
        return None

    # Retorna o texto limpo
    return texto


def normalizar_cabecalho(
    cabecalho: Any,
    separador: str = SEPARADOR_CABECALHO_PADRAO,
    remover_acentos: bool = True,
    converter_minusculas: bool = True
) -> str:
    """
    Normaliza um cabeçalho para uso em DataFrames e automações.

    Args:
        cabecalho (Any):
            Valor utilizado como cabeçalho.
        separador (str):
            Texto usado para substituir espaços e símbolos.
        remover_acentos (bool):
            Define se os acentos serão removidos.
        converter_minusculas (bool):
            Define se o resultado será convertido para minúsculas.

    Returns:
        str:
            Cabeçalho normalizado.

    Raises:
        ValueError:
            Caso o separador esteja vazio.
    """
    # Valida o separador e as opções
    separador_validado = _validar_texto(
        separador,
        "separador"
    )
    _validar_booleano(
        remover_acentos,
        "remover_acentos"
    )
    _validar_booleano(
        converter_minusculas,
        "converter_minusculas"
    )

    # Converte o cabeçalho para texto e aplica a limpeza inicial
    texto = limpar_texto(
        str(cabecalho),
        remover_acentos=remover_acentos,
        converter_minusculas=converter_minusculas
    )

    # Substitui caracteres diferentes de letras e números pelo separador
    texto = re.sub(
        r"[^\w]+",
        separador_validado,
        texto,
        flags=re.UNICODE
    )

    # Reduz separadores consecutivos
    texto = re.sub(
        rf"{re.escape(separador_validado)}+",
        separador_validado,
        texto
    )

    # Remove separadores das extremidades
    texto = texto.strip(separador_validado)

    # Retorna um nome padrão quando o resultado estiver vazio
    return texto or "coluna"


# ----------------------------------------------------------------------------
# FUNÇÕES DE LIMPEZA DE DATAFRAMES
# ----------------------------------------------------------------------------

def limpar_espacos_dataframe(
    dataframe: pd.DataFrame,
    colunas: Iterable[str] | None = None,
    vazio_como_none: bool = False,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Remove espaços e quebras de linha das colunas textuais do DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será limpo.
        colunas (Iterable[str] | None):
            Colunas selecionadas ou None para todas as colunas.
        vazio_como_none (bool):
            Define se textos vazios serão convertidos para None.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame limpo.

    Raises:
        ColunaLimpezaNaoEncontradaError:
            Caso alguma coluna selecionada não exista.
    """
    # Valida o DataFrame e as opções
    _validar_dataframe(dataframe)
    _validar_booleano(
        vazio_como_none,
        "vazio_como_none"
    )
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Define o DataFrame de trabalho
    resultado = dataframe.copy(deep=True) if copiar else dataframe

    # Valida as colunas selecionadas
    selecionadas = _validar_colunas_dataframe(
        resultado,
        colunas
    )

    # Limpa os valores textuais de cada coluna
    for coluna in selecionadas:
        resultado[coluna] = resultado[coluna].map(
            lambda valor: limpar_texto(
                valor,
                vazio_como_none=vazio_como_none
            )
        )

    # Retorna o DataFrame limpo
    return resultado


def normalizar_cabecalhos_dataframe(
    dataframe: pd.DataFrame,
    separador: str = SEPARADOR_CABECALHO_PADRAO,
    garantir_unicos: bool = True,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Normaliza os nomes das colunas de um DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame cujos cabeçalhos serão normalizados.
        separador (str):
            Separador usado no resultado.
        garantir_unicos (bool):
            Define se nomes repetidos receberão sufixos numéricos.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame com os cabeçalhos normalizados.

    Raises:
        ConfiguracaoLimpezaInvalidaError:
            Caso ocorram duplicidades e garantir_unicos seja False.
    """
    # Valida o DataFrame e as opções
    _validar_dataframe(dataframe)
    _validar_booleano(
        garantir_unicos,
        "garantir_unicos"
    )
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Define o DataFrame de trabalho
    resultado = dataframe.copy(deep=True) if copiar else dataframe

    # Cria os nomes normalizados
    normalizados = [
        normalizar_cabecalho(
            coluna,
            separador=separador
        )
        for coluna in resultado.columns
    ]

    # Trata nomes duplicados quando solicitado
    if garantir_unicos:
        contadores: dict[str, int] = {}
        unicos: list[str] = []

        for nome in normalizados:
            contadores[nome] = contadores.get(nome, 0) + 1
            quantidade = contadores[nome]
            unicos.append(
                nome if quantidade == 1 else f"{nome}{separador}{quantidade}"
            )

        normalizados = unicos
    elif len(normalizados) != len(set(normalizados)):
        raise ConfiguracaoLimpezaInvalidaError(
            "A normalização gerou cabeçalhos duplicados."
        )

    # Aplica os novos nomes
    resultado.columns = normalizados

    # Retorna o DataFrame atualizado
    return resultado


def substituir_valores_dataframe(
    dataframe: pd.DataFrame,
    substituicoes: Mapping[Any, Any],
    colunas: Iterable[str] | None = None,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Substitui valores em colunas selecionadas do DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será alterado.
        substituicoes (Mapping[Any, Any]):
            Mapeamento entre valores atuais e novos valores.
        colunas (Iterable[str] | None):
            Colunas selecionadas ou None para todas as colunas.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame com os valores substituídos.

    Raises:
        TypeError:
            Caso substituicoes não seja um mapeamento.
        ValueError:
            Caso o mapeamento esteja vazio.
    """
    # Valida o DataFrame e as configurações
    _validar_dataframe(dataframe)
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Valida o mapeamento de substituições
    if not isinstance(substituicoes, Mapping):
        raise TypeError(
            "O parâmetro 'substituicoes' deve ser um mapeamento."
        )

    if not substituicoes:
        raise ValueError(
            "O parâmetro 'substituicoes' não pode estar vazio."
        )

    # Define o DataFrame e as colunas de trabalho
    resultado = dataframe.copy(deep=True) if copiar else dataframe
    selecionadas = _validar_colunas_dataframe(
        resultado,
        colunas
    )

    # Substitui os valores somente nas colunas selecionadas
    resultado.loc[:, selecionadas] = resultado.loc[:, selecionadas].replace(
        dict(substituicoes)
    )

    # Retorna o DataFrame alterado
    return resultado


def preencher_valores_ausentes(
    dataframe: pd.DataFrame,
    valor: Any = None,
    valores_por_coluna: Mapping[str, Any] | None = None,
    colunas: Iterable[str] | None = None,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Preenche valores ausentes utilizando valor geral ou por coluna.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será alterado.
        valor (Any):
            Valor geral utilizado no preenchimento.
        valores_por_coluna (Mapping[str, Any] | None):
            Mapeamento opcional de valores específicos por coluna.
        colunas (Iterable[str] | None):
            Colunas utilizadas com o valor geral.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame com os valores ausentes preenchidos.

    Raises:
        ConfiguracaoLimpezaInvalidaError:
            Caso nenhum valor de preenchimento seja informado.
    """
    # Valida o DataFrame e a opção de cópia
    _validar_dataframe(dataframe)
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Impede uma configuração sem valor de preenchimento
    if valor is None and valores_por_coluna is None:
        raise ConfiguracaoLimpezaInvalidaError(
            "Informe 'valor' ou 'valores_por_coluna' para o preenchimento."
        )

    # Define o DataFrame de trabalho
    resultado = dataframe.copy(deep=True) if copiar else dataframe

    # Aplica os valores específicos por coluna
    if valores_por_coluna is not None:
        if not isinstance(valores_por_coluna, Mapping):
            raise TypeError(
                "O parâmetro 'valores_por_coluna' deve ser um mapeamento ou None."
            )

        # Valida as colunas do mapeamento
        _validar_colunas_dataframe(
            resultado,
            valores_por_coluna.keys()
        )

        # Preenche os valores conforme o mapeamento
        resultado = resultado.fillna(
            value=dict(valores_por_coluna)
        )

    # Aplica o valor geral nas colunas selecionadas
    if valor is not None:
        selecionadas = _validar_colunas_dataframe(
            resultado,
            colunas
        )
        resultado.loc[:, selecionadas] = resultado.loc[:, selecionadas].fillna(
            valor
        )

    # Retorna o DataFrame preenchido
    return resultado


def remover_linhas_vazias_dataframe(
    dataframe: pd.DataFrame,
    colunas: Iterable[str] | None = None,
    remover_quando: str = "todas",
    redefinir_indice: bool = True,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Remove linhas vazias de um DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será limpo.
        colunas (Iterable[str] | None):
            Colunas consideradas na identificação das linhas vazias.
        remover_quando (str):
            Regra de remoção: todas ou qualquer.
        redefinir_indice (bool):
            Define se o índice será recriado após a remoção.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame sem as linhas vazias.

    Raises:
        ConfiguracaoLimpezaInvalidaError:
            Caso remover_quando não seja todas ou qualquer.
    """
    # Valida o DataFrame e as opções
    _validar_dataframe(dataframe)
    regra = _validar_texto(
        remover_quando,
        "remover_quando"
    ).lower()
    _validar_booleano(
        redefinir_indice,
        "redefinir_indice"
    )
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Converte a regra para o padrão do pandas
    regras = {
        "todas": "all",
        "qualquer": "any"
    }

    if regra not in regras:
        raise ConfiguracaoLimpezaInvalidaError(
            "O parâmetro 'remover_quando' deve ser 'todas' ou 'qualquer'."
        )

    # Define o DataFrame e as colunas analisadas
    resultado = dataframe.copy(deep=True) if copiar else dataframe
    selecionadas = _validar_colunas_dataframe(
        resultado,
        colunas
    )

    # Remove as linhas de acordo com a regra
    resultado = resultado.dropna(
        axis=0,
        how=regras[regra],
        subset=selecionadas
    )

    # Redefine o índice quando solicitado
    if redefinir_indice:
        resultado = resultado.reset_index(drop=True)

    # Retorna o DataFrame limpo
    return resultado


def remover_colunas_vazias_dataframe(
    dataframe: pd.DataFrame,
    remover_quando: str = "todas",
    copiar: bool = True
) -> pd.DataFrame:
    """
    Remove colunas vazias de um DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será limpo.
        remover_quando (str):
            Regra de remoção: todas ou qualquer.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame sem as colunas vazias.

    Raises:
        ConfiguracaoLimpezaInvalidaError:
            Caso remover_quando não seja todas ou qualquer.
    """
    # Valida o DataFrame e as opções
    _validar_dataframe(dataframe)
    regra = _validar_texto(
        remover_quando,
        "remover_quando"
    ).lower()
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Converte a regra para o padrão do pandas
    regras = {
        "todas": "all",
        "qualquer": "any"
    }

    if regra not in regras:
        raise ConfiguracaoLimpezaInvalidaError(
            "O parâmetro 'remover_quando' deve ser 'todas' ou 'qualquer'."
        )

    # Define o DataFrame de trabalho
    resultado = dataframe.copy(deep=True) if copiar else dataframe

    # Remove as colunas conforme a regra
    return resultado.dropna(
        axis=1,
        how=regras[regra]
    )


def remover_duplicidades_dataframe(
    dataframe: pd.DataFrame,
    colunas: Iterable[str] | None = None,
    manter: str | bool = "primeiro",
    redefinir_indice: bool = True,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Remove registros duplicados de um DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será limpo.
        colunas (Iterable[str] | None):
            Colunas consideradas na identificação das duplicidades.
        manter (str | bool):
            Define se será mantido o primeiro, último ou nenhum registro.
        redefinir_indice (bool):
            Define se o índice será recriado.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame sem os registros duplicados.

    Raises:
        ConfiguracaoLimpezaInvalidaError:
            Caso manter possua valor inválido.
    """
    # Valida o DataFrame e as opções
    _validar_dataframe(dataframe)
    _validar_booleano(
        redefinir_indice,
        "redefinir_indice"
    )
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Converte as opções textuais para o padrão do pandas
    opcoes = {
        "primeiro": "first",
        "ultimo": "last",
        "nenhum": False
    }

    if isinstance(manter, str):
        chave = manter.strip().lower()
        if chave not in opcoes:
            raise ConfiguracaoLimpezaInvalidaError(
                "O parâmetro 'manter' deve ser 'primeiro', 'ultimo', "
                "'nenhum' ou False."
            )
        manter_pandas: str | bool = opcoes[chave]
    elif manter is False:
        manter_pandas = False
    else:
        raise ConfiguracaoLimpezaInvalidaError(
            "O parâmetro 'manter' possui valor inválido."
        )

    # Define o DataFrame e valida as colunas
    resultado = dataframe.copy(deep=True) if copiar else dataframe
    selecionadas = (
        None
        if colunas is None
        else _validar_colunas_dataframe(resultado, colunas)
    )

    # Remove os registros duplicados
    resultado = resultado.drop_duplicates(
        subset=selecionadas,
        keep=manter_pandas
    )

    # Redefine o índice quando solicitado
    if redefinir_indice:
        resultado = resultado.reset_index(drop=True)

    # Retorna o DataFrame limpo
    return resultado


def remover_linhas_por_valor(
    dataframe: pd.DataFrame,
    coluna: str,
    valores: Iterable[Any],
    inverter: bool = False,
    redefinir_indice: bool = True,
    copiar: bool = True
) -> pd.DataFrame:
    """
    Remove ou mantém linhas conforme valores de uma coluna.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será filtrado.
        coluna (str):
            Coluna utilizada no filtro.
        valores (Iterable[Any]):
            Valores considerados na operação.
        inverter (bool):
            Define se apenas as linhas correspondentes serão mantidas.
        redefinir_indice (bool):
            Define se o índice será recriado.
        copiar (bool):
            Define se a operação será executada em uma cópia.

    Returns:
        pd.DataFrame:
            DataFrame filtrado.

    Raises:
        ValueError:
            Caso a coleção de valores esteja vazia.
    """
    # Valida o DataFrame, a coluna e as opções
    _validar_dataframe(dataframe)
    coluna_validada = _validar_texto(
        coluna,
        "coluna"
    )
    _validar_colunas_dataframe(
        dataframe,
        [coluna_validada]
    )
    _validar_booleano(
        inverter,
        "inverter"
    )
    _validar_booleano(
        redefinir_indice,
        "redefinir_indice"
    )
    _validar_booleano(
        copiar,
        "copiar"
    )

    # Valida a coleção de valores
    if isinstance(valores, (str, bytes)) or not isinstance(valores, Iterable):
        raise TypeError(
            "O parâmetro 'valores' deve ser um iterável não textual."
        )

    lista_valores = list(valores)

    if not lista_valores:
        raise ValueError(
            "O parâmetro 'valores' não pode estar vazio."
        )

    # Define o DataFrame de trabalho
    resultado = dataframe.copy(deep=True) if copiar else dataframe

    # Cria a máscara de correspondência
    mascara = resultado[coluna_validada].isin(lista_valores)

    # Mantém ou remove as correspondências
    resultado = resultado.loc[
        mascara if inverter else ~mascara
    ]

    # Redefine o índice quando solicitado
    if redefinir_indice:
        resultado = resultado.reset_index(drop=True)

    # Retorna o DataFrame filtrado
    return resultado


# ----------------------------------------------------------------------------
# FUNÇÕES DE LIMPEZA DE PLANILHAS
# ----------------------------------------------------------------------------

def limpar_celula(
    planilha: Worksheet,
    referencia: str,
    limpar_valor: bool = True,
    limpar_comentario: bool = True,
    limpar_link: bool = True
) -> Cell:
    """
    Limpa o conteúdo complementar de uma célula do Excel.

    Args:
        planilha (Worksheet):
            Planilha que contém a célula.
        referencia (str):
            Referência da célula que será limpa.
        limpar_valor (bool):
            Define se o valor será removido.
        limpar_comentario (bool):
            Define se o comentário será removido.
        limpar_link (bool):
            Define se o hiperlink será removido.

    Returns:
        Cell:
            Célula alterada pela operação.

    Raises:
        ReferenciaLimpezaInvalidaError:
            Caso a referência não represente uma única célula editável.
    """
    # Valida a planilha e as opções
    _validar_planilha(planilha)
    for nome, opcao in {
        "limpar_valor": limpar_valor,
        "limpar_comentario": limpar_comentario,
        "limpar_link": limpar_link
    }.items():
        _validar_booleano(
            opcao,
            nome
        )

    # Valida a referência como célula única
    coordenada, limites = _normalizar_intervalo(referencia)
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites

    if coluna_inicial != coluna_final or linha_inicial != linha_final:
        raise ReferenciaLimpezaInvalidaError(
            "A referência deve representar uma única célula."
        )

    # Obtém a célula e rejeita células mescladas secundárias
    celula = planilha[coordenada]

    if isinstance(celula, MergedCell):
        raise ReferenciaLimpezaInvalidaError(
            f"A célula '{coordenada}' pertence a um intervalo mesclado."
        )

    # Limpa os componentes selecionados
    if limpar_valor:
        celula.value = None
    if limpar_comentario:
        celula.comment = None
    if limpar_link:
        celula.hyperlink = None

    # Retorna a célula alterada
    return celula


def limpar_intervalo(
    planilha: Worksheet,
    referencia: str,
    limpar_valores: bool = True,
    limpar_comentarios: bool = True,
    limpar_links: bool = True
) -> int:
    """
    Limpa as células editáveis de um intervalo.

    Args:
        planilha (Worksheet):
            Planilha que contém o intervalo.
        referencia (str):
            Referência do intervalo que será limpo.
        limpar_valores (bool):
            Define se os valores serão removidos.
        limpar_comentarios (bool):
            Define se os comentários serão removidos.
        limpar_links (bool):
            Define se os hiperlinks serão removidos.

    Returns:
        int:
            Quantidade de células editáveis processadas.

    Raises:
        ReferenciaLimpezaInvalidaError:
            Caso a referência do intervalo seja inválida.
    """
    # Valida a planilha, o intervalo e as opções
    _validar_planilha(planilha)
    _, limites = _normalizar_intervalo(referencia)

    for nome, opcao in {
        "limpar_valores": limpar_valores,
        "limpar_comentarios": limpar_comentarios,
        "limpar_links": limpar_links
    }.items():
        _validar_booleano(
            opcao,
            nome
        )

    # Extrai os limites do intervalo
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites
    quantidade = 0

    # Percorre as células do intervalo
    for linha in range(linha_inicial, linha_final + 1):
        for coluna in range(coluna_inicial, coluna_final + 1):
            celula = planilha.cell(
                row=linha,
                column=coluna
            )

            # Ignora células mescladas secundárias
            if isinstance(celula, MergedCell):
                continue

            # Limpa os componentes selecionados
            if limpar_valores:
                celula.value = None
            if limpar_comentarios:
                celula.comment = None
            if limpar_links:
                celula.hyperlink = None

            quantidade += 1

    # Retorna a quantidade processada
    return quantidade


def limpar_espacos_planilha(
    planilha: Worksheet,
    referencia: str | None = None,
    vazio_como_none: bool = False
) -> int:
    """
    Remove espaços e quebras dos valores textuais de uma planilha.

    Args:
        planilha (Worksheet):
            Planilha que será limpa.
        referencia (str | None):
            Intervalo selecionado ou None para a região utilizada.
        vazio_como_none (bool):
            Define se textos vazios serão convertidos para None.

    Returns:
        int:
            Quantidade de células textuais alteradas.

    Raises:
        ReferenciaLimpezaInvalidaError:
            Caso a referência seja inválida.
    """
    # Valida a planilha e a opção
    _validar_planilha(planilha)
    _validar_booleano(
        vazio_como_none,
        "vazio_como_none"
    )

    # Define os limites analisados
    if referencia is None:
        limites = (1, 1, planilha.max_column, planilha.max_row)
    else:
        _, limites = _normalizar_intervalo(referencia)

    coluna_inicial, linha_inicial, coluna_final, linha_final = limites
    quantidade = 0

    # Percorre as células selecionadas
    for linha in range(linha_inicial, linha_final + 1):
        for coluna in range(coluna_inicial, coluna_final + 1):
            celula = planilha.cell(
                row=linha,
                column=coluna
            )

            # Processa somente valores textuais comuns
            if isinstance(celula.value, str) and not celula.value.startswith("="):
                valor_anterior = celula.value
                valor_novo = limpar_texto(
                    valor_anterior,
                    vazio_como_none=vazio_como_none
                )

                if valor_novo != valor_anterior:
                    celula.value = valor_novo
                    quantidade += 1

    # Retorna a quantidade alterada
    return quantidade


def substituir_valores_planilha(
    planilha: Worksheet,
    substituicoes: Mapping[Any, Any],
    referencia: str | None = None
) -> int:
    """
    Substitui valores exatos em uma planilha ou intervalo.

    Args:
        planilha (Worksheet):
            Planilha que será alterada.
        substituicoes (Mapping[Any, Any]):
            Mapeamento entre valores atuais e novos valores.
        referencia (str | None):
            Intervalo selecionado ou None para a região utilizada.

    Returns:
        int:
            Quantidade de células alteradas.

    Raises:
        TypeError:
            Caso substituicoes não seja um mapeamento.
        ValueError:
            Caso o mapeamento esteja vazio.
    """
    # Valida a planilha e o mapeamento
    _validar_planilha(planilha)

    if not isinstance(substituicoes, Mapping):
        raise TypeError(
            "O parâmetro 'substituicoes' deve ser um mapeamento."
        )

    if not substituicoes:
        raise ValueError(
            "O parâmetro 'substituicoes' não pode estar vazio."
        )

    # Define os limites analisados
    if referencia is None:
        limites = (1, 1, planilha.max_column, planilha.max_row)
    else:
        _, limites = _normalizar_intervalo(referencia)

    coluna_inicial, linha_inicial, coluna_final, linha_final = limites
    quantidade = 0

    # Percorre as células selecionadas
    for linha in range(linha_inicial, linha_final + 1):
        for coluna in range(coluna_inicial, coluna_final + 1):
            celula = planilha.cell(
                row=linha,
                column=coluna
            )

            # Substitui valores encontrados como chaves do mapeamento
            try:
                encontrado = celula.value in substituicoes
            except TypeError:
                encontrado = False

            if encontrado:
                celula.value = substituicoes[celula.value]
                quantidade += 1

    # Retorna a quantidade alterada
    return quantidade


def remover_linhas_vazias_planilha(
    planilha: Worksheet,
    linha_inicial: int = 1,
    ignorar_espacos: bool = True
) -> int:
    """
    Remove linhas completamente vazias da região utilizada da planilha.

    Args:
        planilha (Worksheet):
            Planilha que será limpa.
        linha_inicial (int):
            Primeira linha que poderá ser removida.
        ignorar_espacos (bool):
            Define se textos contendo apenas espaços serão considerados vazios.

    Returns:
        int:
            Quantidade de linhas removidas.

    Raises:
        ValueError:
            Caso linha_inicial seja menor que um.
    """
    # Valida a planilha, a linha inicial e a opção
    _validar_planilha(planilha)

    if isinstance(linha_inicial, bool) or not isinstance(linha_inicial, int):
        raise TypeError(
            "O parâmetro 'linha_inicial' deve ser inteiro."
        )

    if linha_inicial < 1:
        raise ValueError(
            "O parâmetro 'linha_inicial' deve ser maior ou igual a 1."
        )

    _validar_booleano(
        ignorar_espacos,
        "ignorar_espacos"
    )

    # Inicia o contador de remoções
    removidas = 0

    # Percorre as linhas de baixo para cima
    for linha in range(planilha.max_row, linha_inicial - 1, -1):
        valores = [
            planilha.cell(
                row=linha,
                column=coluna
            ).value
            for coluna in range(1, planilha.max_column + 1)
        ]

        # Identifica os valores considerados vazios
        def vazio(valor: Any) -> bool:
            if valor is None:
                return True
            if ignorar_espacos and isinstance(valor, str) and not valor.strip():
                return True
            return False

        # Remove a linha quando todos os valores estão vazios
        if all(vazio(valor) for valor in valores):
            planilha.delete_rows(
                linha,
                1
            )
            removidas += 1

    # Retorna a quantidade removida
    return removidas


def remover_colunas_vazias_planilha(
    planilha: Worksheet,
    coluna_inicial: int = 1,
    ignorar_espacos: bool = True
) -> int:
    """
    Remove colunas completamente vazias da região utilizada da planilha.

    Args:
        planilha (Worksheet):
            Planilha que será limpa.
        coluna_inicial (int):
            Primeira coluna que poderá ser removida.
        ignorar_espacos (bool):
            Define se textos contendo apenas espaços serão considerados vazios.

    Returns:
        int:
            Quantidade de colunas removidas.

    Raises:
        ValueError:
            Caso coluna_inicial seja menor que um.
    """
    # Valida a planilha, a coluna inicial e a opção
    _validar_planilha(planilha)

    if isinstance(coluna_inicial, bool) or not isinstance(coluna_inicial, int):
        raise TypeError(
            "O parâmetro 'coluna_inicial' deve ser inteiro."
        )

    if coluna_inicial < 1:
        raise ValueError(
            "O parâmetro 'coluna_inicial' deve ser maior ou igual a 1."
        )

    _validar_booleano(
        ignorar_espacos,
        "ignorar_espacos"
    )

    # Inicia o contador de remoções
    removidas = 0

    # Percorre as colunas da direita para a esquerda
    for coluna in range(planilha.max_column, coluna_inicial - 1, -1):
        valores = [
            planilha.cell(
                row=linha,
                column=coluna
            ).value
            for linha in range(1, planilha.max_row + 1)
        ]

        # Identifica os valores considerados vazios
        def vazio(valor: Any) -> bool:
            if valor is None:
                return True
            if ignorar_espacos and isinstance(valor, str) and not valor.strip():
                return True
            return False

        # Remove a coluna quando todos os valores estão vazios
        if all(vazio(valor) for valor in valores):
            planilha.delete_cols(
                coluna,
                1
            )
            removidas += 1

    # Retorna a quantidade removida
    return removidas


def remover_linhas_duplicadas_planilha(
    planilha: Worksheet,
    linha_cabecalho: int = 1,
    colunas: Iterable[int] | None = None
) -> int:
    """
    Remove linhas duplicadas de uma planilha preservando a primeira ocorrência.

    Args:
        planilha (Worksheet):
            Planilha que será limpa.
        linha_cabecalho (int):
            Linha do cabeçalho; os dados começam na linha seguinte.
        colunas (Iterable[int] | None):
            Índices das colunas consideradas ou None para todas.

    Returns:
        int:
            Quantidade de linhas duplicadas removidas.

    Raises:
        ValueError:
            Caso algum índice de coluna seja menor que um.
    """
    # Valida a planilha e a linha de cabeçalho
    _validar_planilha(planilha)

    if isinstance(linha_cabecalho, bool) or not isinstance(linha_cabecalho, int):
        raise TypeError(
            "O parâmetro 'linha_cabecalho' deve ser inteiro."
        )

    if linha_cabecalho < 0:
        raise ValueError(
            "O parâmetro 'linha_cabecalho' deve ser maior ou igual a 0."
        )

    # Define as colunas consideradas
    if colunas is None:
        indices = list(range(1, planilha.max_column + 1))
    else:
        if isinstance(colunas, (str, bytes)) or not isinstance(colunas, Iterable):
            raise TypeError(
                "O parâmetro 'colunas' deve ser um iterável não textual ou None."
            )

        indices = list(colunas)

        if not indices:
            raise ValueError(
                "O parâmetro 'colunas' não pode estar vazio."
            )

        if any(
            isinstance(indice, bool)
            or not isinstance(indice, int)
            or indice < 1
            for indice in indices
        ):
            raise ValueError(
                "Todos os índices de colunas devem ser inteiros maiores que zero."
            )

    # Armazena as chaves já encontradas
    encontrados: set[tuple[Any, ...]] = set()
    linhas_remover: list[int] = []

    # Percorre as linhas de dados
    for linha in range(linha_cabecalho + 1, planilha.max_row + 1):
        chave = tuple(
            planilha.cell(
                row=linha,
                column=coluna
            ).value
            for coluna in indices
        )

        # Marca as ocorrências posteriores para remoção
        if chave in encontrados:
            linhas_remover.append(linha)
        else:
            encontrados.add(chave)

    # Remove as linhas de baixo para cima
    for linha in reversed(linhas_remover):
        planilha.delete_rows(
            linha,
            1
        )

    # Retorna a quantidade removida
    return len(linhas_remover)


def limpar_planilha(
    planilha: Worksheet,
    limpar_espacos: bool = True,
    remover_linhas_vazias: bool = True,
    remover_colunas_vazias: bool = True,
    remover_duplicidades: bool = False,
    linha_cabecalho: int = 1
) -> dict[str, int]:
    """
    Executa uma sequência padronizada de limpeza em uma planilha.

    Args:
        planilha (Worksheet):
            Planilha que será limpa.
        limpar_espacos (bool):
            Define se textos serão normalizados.
        remover_linhas_vazias (bool):
            Define se linhas completamente vazias serão removidas.
        remover_colunas_vazias (bool):
            Define se colunas completamente vazias serão removidas.
        remover_duplicidades (bool):
            Define se linhas duplicadas serão removidas.
        linha_cabecalho (int):
            Linha do cabeçalho utilizada na remoção de duplicidades.

    Returns:
        dict[str, int]:
            Resumo com as quantidades processadas em cada etapa.

    Raises:
        TypeError:
            Caso alguma opção possua tipo incompatível.
        ValueError:
            Caso a linha de cabeçalho seja inválida.
    """
    # Valida a planilha e as opções booleanas
    _validar_planilha(planilha)

    for nome, opcao in {
        "limpar_espacos": limpar_espacos,
        "remover_linhas_vazias": remover_linhas_vazias,
        "remover_colunas_vazias": remover_colunas_vazias,
        "remover_duplicidades": remover_duplicidades
    }.items():
        _validar_booleano(
            opcao,
            nome
        )

    # Valida a linha de cabeçalho
    if isinstance(linha_cabecalho, bool) or not isinstance(linha_cabecalho, int):
        raise TypeError(
            "O parâmetro 'linha_cabecalho' deve ser inteiro."
        )

    if linha_cabecalho < 0:
        raise ValueError(
            "O parâmetro 'linha_cabecalho' deve ser maior ou igual a 0."
        )

    # Cria o resumo inicial
    resumo = {
        "celulas_textuais_alteradas": 0,
        "linhas_vazias_removidas": 0,
        "colunas_vazias_removidas": 0,
        "linhas_duplicadas_removidas": 0
    }

    # Normaliza os textos da planilha
    if limpar_espacos:
        resumo["celulas_textuais_alteradas"] = limpar_espacos_planilha(
            planilha
        )

    # Remove as linhas vazias
    if remover_linhas_vazias:
        resumo["linhas_vazias_removidas"] = remover_linhas_vazias_planilha(
            planilha,
            linha_inicial=max(linha_cabecalho + 1, 1)
        )

    # Remove as colunas vazias
    if remover_colunas_vazias:
        resumo["colunas_vazias_removidas"] = remover_colunas_vazias_planilha(
            planilha
        )

    # Remove as linhas duplicadas
    if remover_duplicidades:
        resumo["linhas_duplicadas_removidas"] = remover_linhas_duplicadas_planilha(
            planilha,
            linha_cabecalho=linha_cabecalho
        )

    # Retorna o resumo da limpeza
    return resumo


# ----------------------------------------------------------------------------
# EXPORTAÇÕES PÚBLICAS
# ----------------------------------------------------------------------------

__all__ = [
    "PADRAO_ESPACOS",
    "SEPARADOR_CABECALHO_PADRAO",
    "TEXTOS_AUSENTES_PADRAO",
    "ExcelLimpezaError",
    "DataFrameLimpezaInvalidoError",
    "PlanilhaLimpezaInvalidaError",
    "ReferenciaLimpezaInvalidaError",
    "ColunaLimpezaNaoEncontradaError",
    "ConfiguracaoLimpezaInvalidaError",
    "limpar_texto",
    "normalizar_cabecalho",
    "limpar_espacos_dataframe",
    "normalizar_cabecalhos_dataframe",
    "substituir_valores_dataframe",
    "preencher_valores_ausentes",
    "remover_linhas_vazias_dataframe",
    "remover_colunas_vazias_dataframe",
    "remover_duplicidades_dataframe",
    "remover_linhas_por_valor",
    "limpar_celula",
    "limpar_intervalo",
    "limpar_espacos_planilha",
    "substituir_valores_planilha",
    "remover_linhas_vazias_planilha",
    "remover_colunas_vazias_planilha",
    "remover_linhas_duplicadas_planilha",
    "limpar_planilha"
]
