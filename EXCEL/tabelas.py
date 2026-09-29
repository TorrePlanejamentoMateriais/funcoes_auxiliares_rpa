"""
============================================================
Modulo: EXCEL / tabelas.py
Autor: Tiago Eneas Antunes (Z701038)
Data de Criacao: 29/09/2026
Ultima Alteracao: 29/09/2026
Versao: 1.0.0
Descricao:
    Biblioteca de funcoes auxiliares para criar, consultar, alterar,
    redimensionar, estilizar, renomear, copiar e remover tabelas
    estruturadas em planilhas Excel utilizando openpyxl.
Funcoes disponiveis:
    - listar_tabelas()
    - contar_tabelas()
    - tabela_existe()
    - obter_tabela()
    - obter_planilha_tabela()
    - obter_intervalo_tabela()
    - obter_cabecalhos_tabela()
    - obter_dados_tabela()
    - obter_informacoes_tabela()
    - criar_tabela()
    - renomear_tabela()
    - redimensionar_tabela()
    - definir_estilo_tabela()
    - definir_totais_tabela()
    - adicionar_linha_tabela()
    - adicionar_linhas_tabela()
    - copiar_tabela()
    - remover_tabela()
    - remover_tabelas()
Dependencias:
    - openpyxl
Historico:
    v1.0.0 - 29/09/2026
        - Criacao inicial do modulo.
        - Inclusao das operacoes de consulta de tabelas estruturadas.
        - Inclusao da criacao, renomeacao e redimensionamento.
        - Inclusao de estilos, totais e adicao de registros.
        - Inclusao da copia e remocao de tabelas.
============================================================
"""

# ----------------------------------------------------------------------------
# IMPORTACOES
# ----------------------------------------------------------------------------

# Importa deepcopy para copiar valores compostos sem referencias compartilhadas
from copy import deepcopy

# Importa Iterable e Mapping para validar colecoes e registros
from collections.abc import Iterable, Mapping

# Importa Any para anotacoes de valores flexiveis
from typing import Any

# Importa Workbook para consultas globais de tabelas
from openpyxl import Workbook

# Importa Cell e MergedCell para validar celulas editaveis
from openpyxl.cell.cell import Cell, MergedCell

# Importa utilitarios para interpretar e montar referencias
from openpyxl.utils import get_column_letter, range_boundaries

# Importa classes utilizadas pelas tabelas estruturadas
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

# Importa Worksheet para validacao e manipulacao de planilhas
from openpyxl.worksheet.worksheet import Worksheet


# ----------------------------------------------------------------------------
# CONSTANTES
# ----------------------------------------------------------------------------

# Define o estilo visual padrao das tabelas
ESTILO_TABELA_PADRAO = "TableStyleMedium2"

# Define o prefixo padrao utilizado nos nomes das tabelas
PREFIXO_TABELA_PADRAO = "Tabela"

# Define os valores aceitos para funcoes de total do Excel
FUNCOES_TOTAL_VALIDAS = {
    "average",
    "count",
    "countNums",
    "max",
    "min",
    "stdDev",
    "sum",
    "var"
}


# ----------------------------------------------------------------------------
# EXCECOES
# ----------------------------------------------------------------------------

class ExcelTabelasError(Exception):
    """Erro base para operacoes relacionadas a tabelas Excel."""


class WorkbookTabelasInvalidoError(ExcelTabelasError, TypeError):
    """Indica que o objeto informado nao e um Workbook valido."""


class PlanilhaTabelasInvalidaError(ExcelTabelasError, TypeError):
    """Indica que o objeto informado nao e uma Worksheet valida."""


class TabelaInvalidaError(ExcelTabelasError, TypeError):
    """Indica que o objeto informado nao e uma Table valida."""


class NomeTabelaInvalidoError(ExcelTabelasError, ValueError):
    """Indica que o nome da tabela nao atende as regras do Excel."""


class TabelaNaoEncontradaError(ExcelTabelasError, KeyError):
    """Indica que a tabela solicitada nao foi encontrada."""


class TabelaExistenteError(ExcelTabelasError, FileExistsError):
    """Indica que uma tabela com o mesmo nome ja existe no Workbook."""


class IntervaloTabelaInvalidoError(ExcelTabelasError, ValueError):
    """Indica que o intervalo da tabela e invalido ou insuficiente."""


class CabecalhoTabelaInvalidoError(ExcelTabelasError, ValueError):
    """Indica que os cabecalhos da tabela sao vazios ou duplicados."""


class OperacaoTabelaInvalidaError(ExcelTabelasError, ValueError):
    """Indica que uma operacao nao pode ser executada na tabela."""


# ----------------------------------------------------------------------------
# FUNCOES INTERNAS DE VALIDACAO
# ----------------------------------------------------------------------------

def _validar_booleano(
    valor: bool,
    nome: str
) -> bool:
    """
    Valida estritamente um parametro booleano.

    Args:
        valor (bool):
            Valor que sera validado.
        nome (str):
            Nome utilizado na mensagem de erro.

    Returns:
        bool:
            O proprio valor validado.

    Raises:
        TypeError:
            Caso o valor nao seja booleano.
    """
    if not isinstance(valor, bool):
        raise TypeError(
            f"O parametro '{nome}' deve ser booleano."
        )

    return valor


def _validar_inteiro(
    valor: int,
    nome: str,
    minimo: int | None = None
) -> int:
    """
    Valida um numero inteiro e seu limite minimo opcional.

    Args:
        valor (int):
            Numero inteiro que sera validado.
        nome (str):
            Nome utilizado na mensagem de erro.
        minimo (int | None):
            Menor valor permitido.

    Returns:
        int:
            O proprio valor validado.

    Raises:
        TypeError:
            Caso o valor nao seja inteiro.
        ValueError:
            Caso o valor seja menor que o limite minimo.
    """
    if isinstance(valor, bool) or not isinstance(valor, int):
        raise TypeError(
            f"O parametro '{nome}' deve ser inteiro."
        )

    if minimo is not None and valor < minimo:
        raise ValueError(
            f"O parametro '{nome}' deve ser maior ou igual a {minimo}."
        )

    return valor


def _validar_texto(
    valor: str,
    nome: str
) -> str:
    """
    Valida e normaliza um parametro textual obrigatorio.

    Args:
        valor (str):
            Texto que sera validado.
        nome (str):
            Nome utilizado na mensagem de erro.

    Returns:
        str:
            Texto sem espacos externos.

    Raises:
        TypeError:
            Caso o valor nao seja uma string.
        ValueError:
            Caso o texto esteja vazio.
    """
    if not isinstance(valor, str):
        raise TypeError(
            f"O parametro '{nome}' deve ser uma string."
        )

    texto = valor.strip()

    if not texto:
        raise ValueError(
            f"O parametro '{nome}' nao pode estar vazio."
        )

    return texto


def _validar_workbook(
    workbook: Workbook
) -> Workbook:
    """
    Valida se o objeto informado e um Workbook do openpyxl.

    Args:
        workbook (Workbook):
            Workbook que sera validado.

    Returns:
        Workbook:
            O proprio Workbook validado.

    Raises:
        WorkbookTabelasInvalidoError:
            Caso o objeto nao seja um Workbook.
    """
    if not isinstance(workbook, Workbook):
        raise WorkbookTabelasInvalidoError(
            "O objeto fornecido nao e um Workbook do openpyxl."
        )

    return workbook


def _validar_planilha(
    planilha: Worksheet
) -> Worksheet:
    """
    Valida se o objeto informado e uma Worksheet do openpyxl.

    Args:
        planilha (Worksheet):
            Planilha que sera validada.

    Returns:
        Worksheet:
            A propria planilha validada.

    Raises:
        PlanilhaTabelasInvalidaError:
            Caso o objeto nao seja uma Worksheet.
    """
    if not isinstance(planilha, Worksheet):
        raise PlanilhaTabelasInvalidaError(
            "O objeto fornecido nao e uma planilha do openpyxl."
        )

    return planilha


def _validar_tabela(
    tabela: Table
) -> Table:
    """
    Valida se o objeto informado e uma Table do openpyxl.

    Args:
        tabela (Table):
            Tabela estruturada que sera validada.

    Returns:
        Table:
            A propria tabela validada.

    Raises:
        TabelaInvalidaError:
            Caso o objeto nao seja uma Table.
    """
    if not isinstance(tabela, Table):
        raise TabelaInvalidaError(
            "O objeto fornecido nao e uma tabela do openpyxl."
        )

    return tabela


def _validar_nome_tabela(
    nome: str
) -> str:
    """
    Valida um nome destinado a uma tabela estruturada.

    Args:
        nome (str):
            Nome interno da tabela.

    Returns:
        str:
            Nome validado sem espacos externos.

    Raises:
        NomeTabelaInvalidoError:
            Caso o nome inicie com numero ou contenha simbolos invalidos.
    """
    try:
        nome_validado = _validar_texto(
            nome,
            "nome"
        )
    except (TypeError, ValueError) as erro:
        raise NomeTabelaInvalidoError(
            "O nome da tabela e invalido."
        ) from erro

    if nome_validado[0].isdigit():
        raise NomeTabelaInvalidoError(
            "O nome da tabela nao pode comecar com um numero."
        )

    if not all(
        caractere.isalnum() or caractere in {"_", "."}
        for caractere in nome_validado
    ):
        raise NomeTabelaInvalidoError(
            "O nome da tabela deve possuir somente letras, numeros, "
            "sublinhado ou ponto."
        )

    if nome_validado.upper() in {"R", "C"}:
        raise NomeTabelaInvalidoError(
            "Os nomes 'R' e 'C' sao reservados pelo Excel."
        )

    return nome_validado


def _normalizar_intervalo(
    referencia: str,
    exigir_cabecalho_e_dados: bool = True
) -> tuple[str, tuple[int, int, int, int]]:
    """
    Valida um intervalo e retorna seus limites numericos.

    Args:
        referencia (str):
            Referencia no padrao do Excel.
        exigir_cabecalho_e_dados (bool):
            Define se o intervalo deve possuir pelo menos duas linhas.

    Returns:
        tuple[str, tuple[int, int, int, int]]:
            Referencia normalizada e limites numericos.

    Raises:
        IntervaloTabelaInvalidoError:
            Caso a referencia seja invalida ou insuficiente.
    """
    _validar_booleano(
        exigir_cabecalho_e_dados,
        "exigir_cabecalho_e_dados"
    )

    try:
        intervalo = _validar_texto(
            referencia,
            "referencia"
        ).upper()
        limites = range_boundaries(intervalo)
    except (TypeError, ValueError) as erro:
        raise IntervaloTabelaInvalidoError(
            f"O intervalo informado e invalido: '{referencia}'."
        ) from erro

    if any(limite is None for limite in limites) or min(limites) < 1:
        raise IntervaloTabelaInvalidoError(
            f"O intervalo informado e invalido: '{referencia}'."
        )

    coluna_inicial, linha_inicial, coluna_final, linha_final = limites

    if coluna_inicial > coluna_final or linha_inicial > linha_final:
        raise IntervaloTabelaInvalidoError(
            f"O intervalo informado e invalido: '{referencia}'."
        )

    if exigir_cabecalho_e_dados and linha_final <= linha_inicial:
        raise IntervaloTabelaInvalidoError(
            "O intervalo da tabela deve possuir cabecalho e ao menos uma "
            "linha de dados."
        )

    return intervalo, limites


def _validar_cabecalhos_intervalo(
    planilha: Worksheet,
    limites: tuple[int, int, int, int]
) -> list[str]:
    """
    Valida os cabecalhos presentes na primeira linha de um intervalo.

    Args:
        planilha (Worksheet):
            Planilha que contem os cabecalhos.
        limites (tuple[int, int, int, int]):
            Limites numericos do intervalo.

    Returns:
        list[str]:
            Cabecalhos validados na ordem da tabela.

    Raises:
        CabecalhoTabelaInvalidoError:
            Caso existam cabecalhos vazios ou duplicados.
    """
    _validar_planilha(planilha)
    coluna_inicial, linha_inicial, coluna_final, _ = limites

    cabecalhos: list[str] = []

    for coluna in range(coluna_inicial, coluna_final + 1):
        valor = planilha.cell(
            row=linha_inicial,
            column=coluna
        ).value

        if valor is None or not str(valor).strip():
            raise CabecalhoTabelaInvalidoError(
                "Todos os cabecalhos da tabela devem estar preenchidos."
            )

        cabecalhos.append(str(valor).strip())

    if len(cabecalhos) != len(set(cabecalhos)):
        raise CabecalhoTabelaInvalidoError(
            "Os cabecalhos da tabela nao podem possuir duplicidades."
        )

    return cabecalhos


def _iterar_tabelas_workbook(
    workbook: Workbook
) -> Iterable[tuple[Worksheet, Table]]:
    """
    Percorre as tabelas existentes em todas as planilhas do Workbook.

    Args:
        workbook (Workbook):
            Workbook que sera percorrido.

    Returns:
        Iterable[tuple[Worksheet, Table]]:
            Pares contendo a planilha e a tabela encontrada.

    Raises:
        WorkbookTabelasInvalidoError:
            Caso o objeto nao seja um Workbook.
    """
    _validar_workbook(workbook)

    for planilha in workbook.worksheets:
        for tabela in planilha.tables.values():
            yield planilha, tabela


def _nome_tabela_existe_workbook(
    workbook: Workbook,
    nome: str,
    ignorar_tabela: Table | None = None
) -> bool:
    """
    Verifica se um nome de tabela ja existe no Workbook.

    Args:
        workbook (Workbook):
            Workbook que sera consultado.
        nome (str):
            Nome da tabela procurada.
        ignorar_tabela (Table | None):
            Tabela opcional desconsiderada na busca.

    Returns:
        bool:
            True quando o nome ja esta em uso.

    Raises:
        NomeTabelaInvalidoError:
            Caso o nome seja invalido.
    """
    nome_validado = _validar_nome_tabela(nome)

    return any(
        tabela is not ignorar_tabela
        and tabela.displayName.casefold() == nome_validado.casefold()
        for _, tabela in _iterar_tabelas_workbook(workbook)
    )


def _obter_planilha_da_tabela(
    tabela: Table
) -> Worksheet:
    """
    Localiza a planilha a qual uma tabela pertence.

    Args:
        tabela (Table):
            Tabela que sera localizada.

    Returns:
        Worksheet:
            Planilha que contem a tabela.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao esteja vinculada a uma Worksheet acessivel.
    """
    _validar_tabela(tabela)

    # O openpyxl nao mantem uma referencia publica ao pai da tabela.
    # A funcao utiliza o atributo interno associado durante add_table.
    planilha = getattr(tabela, "_worksheet", None)

    if isinstance(planilha, Worksheet):
        return planilha

    raise TabelaNaoEncontradaError(
        "Nao foi possivel identificar a planilha da tabela informada."
    )


def _resolver_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table
) -> tuple[Worksheet, Table]:
    """
    Resolve uma tabela a partir de Workbook, Worksheet, nome ou objeto.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha utilizada na busca.
        referencia (str | Table):
            Nome ou objeto da tabela.

    Returns:
        tuple[Worksheet, Table]:
            Planilha e tabela localizadas.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada na origem.
    """
    if not isinstance(origem, (Workbook, Worksheet)):
        raise TypeError(
            "O parametro 'origem' deve ser Workbook ou Worksheet."
        )

    if isinstance(referencia, Table):
        if isinstance(origem, Worksheet):
            if referencia in origem.tables.values():
                return origem, referencia
        else:
            for planilha, tabela in _iterar_tabelas_workbook(origem):
                if tabela is referencia:
                    return planilha, tabela

        raise TabelaNaoEncontradaError(
            "A tabela informada nao pertence a origem pesquisada."
        )

    nome = _validar_nome_tabela(referencia)

    if isinstance(origem, Worksheet):
        for tabela in origem.tables.values():
            if tabela.displayName.casefold() == nome.casefold():
                return origem, tabela
    else:
        for planilha, tabela in _iterar_tabelas_workbook(origem):
            if tabela.displayName.casefold() == nome.casefold():
                return planilha, tabela

    raise TabelaNaoEncontradaError(
        f"A tabela '{nome}' nao foi encontrada."
    )


# ----------------------------------------------------------------------------
# FUNCOES DE CONSULTA
# ----------------------------------------------------------------------------

def listar_tabelas(
    origem: Workbook | Worksheet
) -> list[str]:
    """
    Lista os nomes das tabelas existentes em um Workbook ou Worksheet.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que sera consultada.

    Returns:
        list[str]:
            Nomes das tabelas na ordem encontrada.

    Raises:
        TypeError:
            Caso a origem nao seja Workbook ou Worksheet.
    """
    if isinstance(origem, Worksheet):
        return [
            tabela.displayName
            for tabela in origem.tables.values()
        ]

    if isinstance(origem, Workbook):
        return [
            tabela.displayName
            for _, tabela in _iterar_tabelas_workbook(origem)
        ]

    raise TypeError(
        "O parametro 'origem' deve ser Workbook ou Worksheet."
    )


def contar_tabelas(
    origem: Workbook | Worksheet
) -> int:
    """
    Retorna a quantidade de tabelas estruturadas na origem.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que sera consultada.

    Returns:
        int:
            Quantidade de tabelas encontradas.

    Raises:
        TypeError:
            Caso a origem nao seja Workbook ou Worksheet.
    """
    return len(listar_tabelas(origem))


def tabela_existe(
    origem: Workbook | Worksheet,
    nome_tabela: str,
    considerar_maiusculas: bool = False
) -> bool:
    """
    Verifica se uma tabela existe em um Workbook ou Worksheet.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que sera consultada.
        nome_tabela (str):
            Nome da tabela procurada.
        considerar_maiusculas (bool):
            Define se a comparacao diferencia maiusculas e minusculas.

    Returns:
        bool:
            True quando a tabela existe; caso contrario, False.

    Raises:
        NomeTabelaInvalidoError:
            Caso o nome informado seja invalido.
    """
    nome = _validar_nome_tabela(nome_tabela)
    _validar_booleano(
        considerar_maiusculas,
        "considerar_maiusculas"
    )

    nomes = listar_tabelas(origem)

    if considerar_maiusculas:
        return nome in nomes

    return nome.casefold() in {
        item.casefold()
        for item in nomes
    }


def obter_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table
) -> Table:
    """
    Obtem uma tabela por nome ou objeto Table.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha utilizada na consulta.
        referencia (str | Table):
            Nome ou objeto da tabela.

    Returns:
        Table:
            Tabela localizada na origem.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada.
    """
    _, tabela = _resolver_tabela(
        origem,
        referencia
    )

    return tabela


def obter_planilha_tabela(
    workbook: Workbook,
    referencia: str | Table
) -> Worksheet:
    """
    Retorna a planilha que contem uma tabela do Workbook.

    Args:
        workbook (Workbook):
            Workbook utilizado na consulta.
        referencia (str | Table):
            Nome ou objeto da tabela.

    Returns:
        Worksheet:
            Planilha que contem a tabela.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada.
    """
    planilha, _ = _resolver_tabela(
        workbook,
        referencia
    )

    return planilha


def obter_intervalo_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table
) -> str:
    """
    Retorna o intervalo ocupado por uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha utilizada na consulta.
        referencia (str | Table):
            Nome ou objeto da tabela.

    Returns:
        str:
            Referencia do intervalo no padrao do Excel.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada.
    """
    tabela = obter_tabela(
        origem,
        referencia
    )

    return tabela.ref


def obter_cabecalhos_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table
) -> list[str]:
    """
    Retorna os cabecalhos atuais de uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha utilizada na consulta.
        referencia (str | Table):
            Nome ou objeto da tabela.

    Returns:
        list[str]:
            Cabecalhos na ordem da tabela.

    Raises:
        CabecalhoTabelaInvalidoError:
            Caso os cabecalhos do intervalo sejam invalidos.
    """
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    _, limites = _normalizar_intervalo(
        tabela.ref,
        exigir_cabecalho_e_dados=False
    )

    return _validar_cabecalhos_intervalo(
        planilha,
        limites
    )


def obter_dados_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    como_dicionarios: bool = True,
    incluir_linha_totais: bool = False
) -> list[Any]:
    """
    Retorna os registros armazenados em uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha utilizada na consulta.
        referencia (str | Table):
            Nome ou objeto da tabela.
        como_dicionarios (bool):
            Define se cada linha sera retornada como dicionario.
        incluir_linha_totais (bool):
            Define se a linha de totais sera incluida nos dados.

    Returns:
        list[Any]:
            Registros da tabela como dicionarios ou listas.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada.
    """
    _validar_booleano(
        como_dicionarios,
        "como_dicionarios"
    )
    _validar_booleano(
        incluir_linha_totais,
        "incluir_linha_totais"
    )

    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    cabecalhos = obter_cabecalhos_tabela(
        planilha,
        tabela
    )
    _, limites = _normalizar_intervalo(
        tabela.ref,
        exigir_cabecalho_e_dados=False
    )
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites

    if tabela.totalsRowShown and not incluir_linha_totais:
        linha_final -= 1

    registros: list[Any] = []

    for linha in range(linha_inicial + 1, linha_final + 1):
        valores = [
            planilha.cell(
                row=linha,
                column=coluna
            ).value
            for coluna in range(coluna_inicial, coluna_final + 1)
        ]

        if como_dicionarios:
            registros.append(dict(zip(cabecalhos, valores)))
        else:
            registros.append(valores)

    return registros


def obter_informacoes_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table
) -> dict[str, Any]:
    """
    Retorna informacoes estruturais e visuais de uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha utilizada na consulta.
        referencia (str | Table):
            Nome ou objeto da tabela.

    Returns:
        dict[str, Any]:
            Informacoes da tabela, intervalo, estilo e dimensoes.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada.
    """
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    _, limites = _normalizar_intervalo(
        tabela.ref,
        exigir_cabecalho_e_dados=False
    )
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites
    estilo = tabela.tableStyleInfo

    return {
        "nome": tabela.displayName,
        "planilha": planilha.title,
        "intervalo": tabela.ref,
        "cabecalhos": obter_cabecalhos_tabela(planilha, tabela),
        "quantidade_colunas": coluna_final - coluna_inicial + 1,
        "quantidade_linhas_dados": max(
            linha_final - linha_inicial - (1 if tabela.totalsRowShown else 0),
            0
        ),
        "linha_totais": bool(tabela.totalsRowShown),
        "estilo": estilo.name if estilo is not None else None,
        "linhas_alternadas": (
            estilo.showRowStripes
            if estilo is not None
            else None
        ),
        "colunas_alternadas": (
            estilo.showColumnStripes
            if estilo is not None
            else None
        ),
        "primeira_coluna_destacada": (
            estilo.showFirstColumn
            if estilo is not None
            else None
        ),
        "ultima_coluna_destacada": (
            estilo.showLastColumn
            if estilo is not None
            else None
        )
    }


# ----------------------------------------------------------------------------
# FUNCOES DE CRIACAO E ALTERACAO
# ----------------------------------------------------------------------------

def criar_tabela(
    planilha: Worksheet,
    nome_tabela: str,
    referencia: str | None = None,
    estilo: str = ESTILO_TABELA_PADRAO,
    linhas_alternadas: bool = True,
    colunas_alternadas: bool = False,
    primeira_coluna_destacada: bool = False,
    ultima_coluna_destacada: bool = False
) -> Table:
    """
    Cria uma tabela estruturada em uma planilha.

    Args:
        planilha (Worksheet):
            Planilha que recebera a tabela.
        nome_tabela (str):
            Nome interno da tabela.
        referencia (str | None):
            Intervalo da tabela ou None para a regiao utilizada.
        estilo (str):
            Nome do estilo visual da tabela.
        linhas_alternadas (bool):
            Define a exibicao de linhas alternadas.
        colunas_alternadas (bool):
            Define a exibicao de colunas alternadas.
        primeira_coluna_destacada (bool):
            Define o destaque da primeira coluna.
        ultima_coluna_destacada (bool):
            Define o destaque da ultima coluna.

    Returns:
        Table:
            Tabela criada e adicionada a planilha.

    Raises:
        TabelaExistenteError:
            Caso o nome ja exista no Workbook.
        CabecalhoTabelaInvalidoError:
            Caso os cabecalhos sejam invalidos.
    """
    _validar_planilha(planilha)
    nome = _validar_nome_tabela(nome_tabela)
    estilo_validado = _validar_texto(
        estilo,
        "estilo"
    )

    for nome_opcao, valor in {
        "linhas_alternadas": linhas_alternadas,
        "colunas_alternadas": colunas_alternadas,
        "primeira_coluna_destacada": primeira_coluna_destacada,
        "ultima_coluna_destacada": ultima_coluna_destacada
    }.items():
        _validar_booleano(
            valor,
            nome_opcao
        )

    workbook = planilha.parent

    if _nome_tabela_existe_workbook(workbook, nome):
        raise TabelaExistenteError(
            f"A tabela '{nome}' ja existe no Workbook."
        )

    if referencia is None:
        referencia_tabela = (
            f"A1:{get_column_letter(planilha.max_column)}{planilha.max_row}"
        )
    else:
        referencia_tabela = referencia

    intervalo, limites = _normalizar_intervalo(referencia_tabela)
    _validar_cabecalhos_intervalo(
        planilha,
        limites
    )

    tabela = Table(
        displayName=nome,
        ref=intervalo
    )
    tabela.tableStyleInfo = TableStyleInfo(
        name=estilo_validado,
        showFirstColumn=primeira_coluna_destacada,
        showLastColumn=ultima_coluna_destacada,
        showRowStripes=linhas_alternadas,
        showColumnStripes=colunas_alternadas
    )

    planilha.add_table(tabela)
    setattr(tabela, "_worksheet", planilha)

    return tabela


def renomear_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    novo_nome: str
) -> Table:
    """
    Renomeia uma tabela estruturada.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela atual.
        novo_nome (str):
            Novo nome interno da tabela.

    Returns:
        Table:
            A propria tabela renomeada.

    Raises:
        TabelaExistenteError:
            Caso o novo nome ja exista no Workbook.
    """
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    nome = _validar_nome_tabela(novo_nome)

    if nome == tabela.displayName:
        return tabela

    workbook = planilha.parent

    if _nome_tabela_existe_workbook(
        workbook,
        nome,
        ignorar_tabela=tabela
    ):
        raise TabelaExistenteError(
            f"A tabela '{nome}' ja existe no Workbook."
        )

    nome_anterior = tabela.displayName
    planilha.tables.pop(nome_anterior)
    tabela.displayName = nome
    tabela.name = nome
    planilha.tables.add(tabela)
    setattr(tabela, "_worksheet", planilha)

    return tabela


def redimensionar_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    novo_intervalo: str
) -> Table:
    """
    Redimensiona o intervalo ocupado por uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela.
        novo_intervalo (str):
            Novo intervalo completo da tabela.

    Returns:
        Table:
            A propria tabela redimensionada.

    Raises:
        CabecalhoTabelaInvalidoError:
            Caso os novos cabecalhos sejam invalidos.
    """
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    intervalo, limites = _normalizar_intervalo(novo_intervalo)
    _validar_cabecalhos_intervalo(
        planilha,
        limites
    )

    tabela.ref = intervalo
    return tabela


def definir_estilo_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    estilo: str = ESTILO_TABELA_PADRAO,
    linhas_alternadas: bool = True,
    colunas_alternadas: bool = False,
    primeira_coluna_destacada: bool = False,
    ultima_coluna_destacada: bool = False
) -> Table:
    """
    Define o estilo visual de uma tabela estruturada.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela.
        estilo (str):
            Nome do estilo visual.
        linhas_alternadas (bool):
            Define a exibicao de linhas alternadas.
        colunas_alternadas (bool):
            Define a exibicao de colunas alternadas.
        primeira_coluna_destacada (bool):
            Define o destaque da primeira coluna.
        ultima_coluna_destacada (bool):
            Define o destaque da ultima coluna.

    Returns:
        Table:
            A propria tabela com o novo estilo.

    Raises:
        ValueError:
            Caso o nome do estilo esteja vazio.
    """
    tabela = obter_tabela(
        origem,
        referencia
    )
    estilo_validado = _validar_texto(
        estilo,
        "estilo"
    )

    for nome_opcao, valor in {
        "linhas_alternadas": linhas_alternadas,
        "colunas_alternadas": colunas_alternadas,
        "primeira_coluna_destacada": primeira_coluna_destacada,
        "ultima_coluna_destacada": ultima_coluna_destacada
    }.items():
        _validar_booleano(
            valor,
            nome_opcao
        )

    tabela.tableStyleInfo = TableStyleInfo(
        name=estilo_validado,
        showFirstColumn=primeira_coluna_destacada,
        showLastColumn=ultima_coluna_destacada,
        showRowStripes=linhas_alternadas,
        showColumnStripes=colunas_alternadas
    )

    return tabela


def definir_totais_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    totais: Mapping[str, str | None],
    rotulo_primeira_coluna: str = "Total"
) -> Table:
    """
    Configura a linha de totais de uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela.
        totais (Mapping[str, str | None]):
            Mapeamento entre cabecalhos e funcoes de total.
        rotulo_primeira_coluna (str):
            Rotulo exibido na primeira coluna da linha de totais.

    Returns:
        Table:
            A propria tabela com totais configurados.

    Raises:
        OperacaoTabelaInvalidaError:
            Caso uma coluna ou funcao de total seja invalida.
    """
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )

    if not isinstance(totais, Mapping) or not totais:
        raise TypeError(
            "O parametro 'totais' deve ser um mapeamento nao vazio."
        )

    rotulo = _validar_texto(
        rotulo_primeira_coluna,
        "rotulo_primeira_coluna"
    )
    cabecalhos = obter_cabecalhos_tabela(
        planilha,
        tabela
    )

    ausentes = [
        coluna
        for coluna in totais
        if coluna not in cabecalhos
    ]

    if ausentes:
        raise OperacaoTabelaInvalidaError(
            f"As colunas nao pertencem a tabela: {ausentes}."
        )

    _, limites = _normalizar_intervalo(
        tabela.ref,
        exigir_cabecalho_e_dados=False
    )
    coluna_inicial, _, coluna_final, linha_final = limites

    if not tabela.totalsRowShown:
        linha_totais = linha_final + 1
        tabela.ref = (
            f"{get_column_letter(coluna_inicial)}{limites[1]}:"
            f"{get_column_letter(coluna_final)}{linha_totais}"
        )
    else:
        linha_totais = linha_final

    tabela.totalsRowShown = True
    tabela.totalsRowCount = 1

    # Garante a criacao das colunas da tabela antes de aplicar os totais.
    if not tabela.tableColumns:
        tabela._initialise_columns()
        for coluna_tabela, cabecalho in zip(tabela.tableColumns, cabecalhos):
            coluna_tabela.name = cabecalho

    for indice, coluna_tabela in enumerate(tabela.tableColumns):
        cabecalho = cabecalhos[indice]
        celula = planilha.cell(
            row=linha_totais,
            column=coluna_inicial + indice
        )

        if indice == 0:
            coluna_tabela.totalsRowLabel = rotulo
            celula.value = rotulo

        if cabecalho in totais:
            funcao = totais[cabecalho]

            if funcao is None:
                coluna_tabela.totalsRowFunction = None
                celula.value = None
            else:
                funcao_validada = _validar_texto(
                    funcao,
                    f"totais[{cabecalho}]"
                )

                if funcao_validada not in FUNCOES_TOTAL_VALIDAS:
                    raise OperacaoTabelaInvalidaError(
                        f"A funcao de total e invalida: '{funcao_validada}'."
                    )

                coluna_tabela.totalsRowFunction = funcao_validada

    return tabela


def adicionar_linha_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    valores: Mapping[str, Any] | Iterable[Any]
) -> list[Cell]:
    """
    Adiciona uma linha de dados ao final de uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela.
        valores (Mapping[str, Any] | Iterable[Any]):
            Registro por cabecalho ou colecao ordenada de valores.

    Returns:
        list[Cell]:
            Celulas preenchidas na nova linha.

    Raises:
        OperacaoTabelaInvalidaError:
            Caso a quantidade ou os nomes dos valores sejam invalidos.
    """
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    cabecalhos = obter_cabecalhos_tabela(
        planilha,
        tabela
    )

    if isinstance(valores, Mapping):
        extras = [
            chave
            for chave in valores
            if chave not in cabecalhos
        ]

        if extras:
            raise OperacaoTabelaInvalidaError(
                f"As chaves nao pertencem a tabela: {extras}."
            )

        dados = [
            valores.get(cabecalho)
            for cabecalho in cabecalhos
        ]
    else:
        if isinstance(valores, (str, bytes)) or not isinstance(
            valores,
            Iterable
        ):
            raise TypeError(
                "O parametro 'valores' deve ser um mapeamento ou iteravel "
                "nao textual."
            )

        dados = list(valores)

        if len(dados) != len(cabecalhos):
            raise OperacaoTabelaInvalidaError(
                "A quantidade de valores deve corresponder a quantidade de "
                "colunas da tabela."
            )

    _, limites = _normalizar_intervalo(
        tabela.ref,
        exigir_cabecalho_e_dados=False
    )
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites

    if tabela.totalsRowShown:
        linha_destino = linha_final
        planilha.insert_rows(linha_destino)
        nova_linha_final = linha_final + 1
    else:
        linha_destino = linha_final + 1
        nova_linha_final = linha_destino

    celulas: list[Cell] = []

    for deslocamento, valor in enumerate(dados):
        celula = planilha.cell(
            row=linha_destino,
            column=coluna_inicial + deslocamento,
            value=deepcopy(valor)
        )
        celulas.append(celula)

    tabela.ref = (
        f"{get_column_letter(coluna_inicial)}{linha_inicial}:"
        f"{get_column_letter(coluna_final)}{nova_linha_final}"
    )

    return celulas


def adicionar_linhas_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    linhas: Iterable[Mapping[str, Any] | Iterable[Any]]
) -> int:
    """
    Adiciona varias linhas ao final de uma tabela.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela.
        linhas (Iterable[Mapping[str, Any] | Iterable[Any]]):
            Colecao de registros adicionados.

    Returns:
        int:
            Quantidade de linhas adicionadas.

    Raises:
        TypeError:
            Caso linhas nao seja um iteravel nao textual.
    """
    if isinstance(linhas, (str, bytes)) or not isinstance(linhas, Iterable):
        raise TypeError(
            "O parametro 'linhas' deve ser um iteravel nao textual."
        )

    registros = list(linhas)

    for valores in registros:
        adicionar_linha_tabela(
            origem,
            referencia,
            valores
        )

    return len(registros)


def copiar_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    planilha_destino: Worksheet,
    celula_inicial: str,
    novo_nome: str,
    copiar_estilo: bool = True
) -> Table:
    """
    Copia os dados de uma tabela para outra planilha.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela de origem.
        referencia (str | Table):
            Nome ou objeto da tabela de origem.
        planilha_destino (Worksheet):
            Planilha que recebera a copia.
        celula_inicial (str):
            Celula superior esquerda da tabela copiada.
        novo_nome (str):
            Nome da nova tabela.
        copiar_estilo (bool):
            Define se o estilo visual sera copiado.

    Returns:
        Table:
            Nova tabela criada na planilha de destino.

    Raises:
        TabelaExistenteError:
            Caso o novo nome ja exista no Workbook de destino.
    """
    planilha_origem, tabela_origem = _resolver_tabela(
        origem,
        referencia
    )
    _validar_planilha(planilha_destino)
    nome = _validar_nome_tabela(novo_nome)
    _validar_booleano(
        copiar_estilo,
        "copiar_estilo"
    )

    if _nome_tabela_existe_workbook(planilha_destino.parent, nome):
        raise TabelaExistenteError(
            f"A tabela '{nome}' ja existe no Workbook de destino."
        )

    celula, limites_destino = _normalizar_intervalo(
        celula_inicial,
        exigir_cabecalho_e_dados=False
    )
    coluna_destino, linha_destino, coluna_final_destino, linha_final_destino = (
        limites_destino
    )

    if coluna_destino != coluna_final_destino or linha_destino != linha_final_destino:
        raise IntervaloTabelaInvalidoError(
            "A celula inicial deve representar uma unica celula."
        )

    _, limites_origem = _normalizar_intervalo(
        tabela_origem.ref,
        exigir_cabecalho_e_dados=False
    )
    coluna_inicial_origem, linha_inicial_origem, coluna_final_origem, linha_final_origem = (
        limites_origem
    )

    for deslocamento_linha, linha in enumerate(
        range(linha_inicial_origem, linha_final_origem + 1)
    ):
        for deslocamento_coluna, coluna in enumerate(
            range(coluna_inicial_origem, coluna_final_origem + 1)
        ):
            origem_celula = planilha_origem.cell(
                row=linha,
                column=coluna
            )
            destino_celula = planilha_destino.cell(
                row=linha_destino + deslocamento_linha,
                column=coluna_destino + deslocamento_coluna,
                value=deepcopy(origem_celula.value)
            )

            if origem_celula.has_style:
                destino_celula._style = deepcopy(origem_celula._style)
                destino_celula.number_format = origem_celula.number_format

    quantidade_colunas = coluna_final_origem - coluna_inicial_origem + 1
    quantidade_linhas = linha_final_origem - linha_inicial_origem + 1
    referencia_destino = (
        f"{get_column_letter(coluna_destino)}{linha_destino}:"
        f"{get_column_letter(coluna_destino + quantidade_colunas - 1)}"
        f"{linha_destino + quantidade_linhas - 1}"
    )

    if copiar_estilo and tabela_origem.tableStyleInfo is not None:
        estilo = tabela_origem.tableStyleInfo
        nova_tabela = criar_tabela(
            planilha_destino,
            nome,
            referencia=referencia_destino,
            estilo=estilo.name,
            linhas_alternadas=bool(estilo.showRowStripes),
            colunas_alternadas=bool(estilo.showColumnStripes),
            primeira_coluna_destacada=bool(estilo.showFirstColumn),
            ultima_coluna_destacada=bool(estilo.showLastColumn)
        )
    else:
        nova_tabela = criar_tabela(
            planilha_destino,
            nome,
            referencia=referencia_destino
        )

    return nova_tabela


# ----------------------------------------------------------------------------
# FUNCOES DE REMOCAO
# ----------------------------------------------------------------------------

def remover_tabela(
    origem: Workbook | Worksheet,
    referencia: str | Table,
    limpar_dados: bool = False
) -> str:
    """
    Remove uma tabela estruturada da planilha.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que contem a tabela.
        referencia (str | Table):
            Nome ou objeto da tabela.
        limpar_dados (bool):
            Define se os valores do intervalo tambem serao removidos.

    Returns:
        str:
            Nome da tabela removida.

    Raises:
        TabelaNaoEncontradaError:
            Caso a tabela nao seja encontrada.
    """
    _validar_booleano(
        limpar_dados,
        "limpar_dados"
    )
    planilha, tabela = _resolver_tabela(
        origem,
        referencia
    )
    nome = tabela.displayName

    if limpar_dados:
        _, limites = _normalizar_intervalo(
            tabela.ref,
            exigir_cabecalho_e_dados=False
        )
        coluna_inicial, linha_inicial, coluna_final, linha_final = limites

        for linha in range(linha_inicial, linha_final + 1):
            for coluna in range(coluna_inicial, coluna_final + 1):
                celula = planilha.cell(
                    row=linha,
                    column=coluna
                )

                if not isinstance(celula, MergedCell):
                    celula.value = None

    planilha.tables.pop(nome)
    return nome


def remover_tabelas(
    origem: Workbook | Worksheet,
    referencias: Iterable[str | Table] | None = None,
    limpar_dados: bool = False,
    ignorar_ausentes: bool = False
) -> list[str]:
    """
    Remove varias tabelas de um Workbook ou Worksheet.

    Args:
        origem (Workbook | Worksheet):
            Workbook ou planilha que sera alterada.
        referencias (Iterable[str | Table] | None):
            Tabelas removidas ou None para remover todas.
        limpar_dados (bool):
            Define se os valores dos intervalos serao removidos.
        ignorar_ausentes (bool):
            Define se referencias inexistentes serao ignoradas.

    Returns:
        list[str]:
            Nomes das tabelas removidas.

    Raises:
        TypeError:
            Caso referencias nao seja um iteravel nao textual ou None.
    """
    _validar_booleano(
        limpar_dados,
        "limpar_dados"
    )
    _validar_booleano(
        ignorar_ausentes,
        "ignorar_ausentes"
    )

    if referencias is None:
        itens: list[str | Table] = list(listar_tabelas(origem))
    else:
        if isinstance(referencias, (str, bytes)) or not isinstance(
            referencias,
            Iterable
        ):
            raise TypeError(
                "O parametro 'referencias' deve ser um iteravel nao textual "
                "ou None."
            )

        itens = list(referencias)

    removidas: list[str] = []

    for referencia in itens:
        try:
            nome = remover_tabela(
                origem,
                referencia,
                limpar_dados=limpar_dados
            )
        except TabelaNaoEncontradaError:
            if ignorar_ausentes:
                continue
            raise

        removidas.append(nome)

    return removidas


# ----------------------------------------------------------------------------
# EXPORTACOES PUBLICAS
# ----------------------------------------------------------------------------

__all__ = [
    "ESTILO_TABELA_PADRAO",
    "PREFIXO_TABELA_PADRAO",
    "FUNCOES_TOTAL_VALIDAS",
    "ExcelTabelasError",
    "WorkbookTabelasInvalidoError",
    "PlanilhaTabelasInvalidaError",
    "TabelaInvalidaError",
    "NomeTabelaInvalidoError",
    "TabelaNaoEncontradaError",
    "TabelaExistenteError",
    "IntervaloTabelaInvalidoError",
    "CabecalhoTabelaInvalidoError",
    "OperacaoTabelaInvalidaError",
    "listar_tabelas",
    "contar_tabelas",
    "tabela_existe",
    "obter_tabela",
    "obter_planilha_tabela",
    "obter_intervalo_tabela",
    "obter_cabecalhos_tabela",
    "obter_dados_tabela",
    "obter_informacoes_tabela",
    "criar_tabela",
    "renomear_tabela",
    "redimensionar_tabela",
    "definir_estilo_tabela",
    "definir_totais_tabela",
    "adicionar_linha_tabela",
    "adicionar_linhas_tabela",
    "copiar_tabela",
    "remover_tabela",
    "remover_tabelas"
]
