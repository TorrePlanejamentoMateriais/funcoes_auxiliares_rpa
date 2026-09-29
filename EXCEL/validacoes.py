"""
============================================================
Módulo: EXCEL / validacoes.py
Autor: Tiago Eneas Antunes (Z701038)
Data de Criação: 29/09/2026
Última Alteração: 29/09/2026
Versão: 1.0.0
Descrição:
    Biblioteca de funções auxiliares para validar arquivos, extensões,
    Workbooks, planilhas, células, intervalos, linhas, colunas,
    DataFrames, nomes de planilhas, cores e parâmetros utilizados pelos
    módulos da biblioteca de automação Excel.
Funções disponíveis:
    - validar_booleano()
    - validar_inteiro()
    - validar_numero()
    - validar_texto()
    - validar_iteravel()
    - validar_mapeamento()
    - normalizar_caminho()
    - validar_extensao_excel()
    - validar_arquivo_excel()
    - validar_diretorio()
    - validar_workbook()
    - validar_planilha()
    - validar_nome_planilha()
    - validar_planilha_existente()
    - validar_planilha_inexistente()
    - validar_referencia_celula()
    - validar_referencia_intervalo()
    - validar_linha()
    - validar_coluna()
    - validar_dataframe()
    - validar_colunas_dataframe()
    - validar_cor_hexadecimal()
    - validar_formato_numero()
    - validar_nome_tabela()
Dependências:
    - pandas
    - openpyxl
    - pathlib
Histórico:
    v1.0.0 - 29/09/2026
        - Criação inicial do módulo.
        - Inclusão das validações de tipos e parâmetros básicos.
        - Inclusão das validações de arquivos e diretórios.
        - Inclusão das validações de Workbook e planilhas.
        - Inclusão das validações de células, intervalos e dimensões.
        - Inclusão das validações de DataFrames, cores e tabelas.
============================================================
"""

# ----------------------------------------------------------------------------
# IMPORTAÇÕES
# ----------------------------------------------------------------------------

# Importa Path para normalização e validação de caminhos
from pathlib import Path

# Importa Iterable e Mapping para validar coleções e mapeamentos
from collections.abc import Iterable, Mapping

# Importa Any para aceitar diferentes tipos de valores
from typing import Any

# Importa pandas para validar DataFrames
import pandas as pd

# Importa MergedCell para identificar células mescladas secundárias
from openpyxl.cell.cell import MergedCell

# Importa utilitários para interpretar referências e colunas
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

# Importa Workbook e Worksheet para validar objetos do openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


# ----------------------------------------------------------------------------
# CONSTANTES
# ----------------------------------------------------------------------------

# Define as extensões reconhecidas como arquivos Excel
EXTENSOES_EXCEL = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
    ".xls"
}

# Define as extensões manipuladas diretamente pelo openpyxl
EXTENSOES_OPENPYXL = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm"
}

# Define os limites estruturais das planilhas modernas
LIMITE_LINHAS_EXCEL = 1_048_576
LIMITE_COLUNAS_EXCEL = 16_384
LIMITE_CARACTERES_PLANILHA = 31
LIMITE_CARACTERES_CELULA = 32_767
ULTIMA_COLUNA_EXCEL = "XFD"

# Define os caracteres proibidos nos nomes das planilhas
CARACTERES_INVALIDOS_PLANILHA = {
    "[",
    "]",
    ":",
    "*",
    "?",
    "/",
    "\\"
}

# Define os nomes reservados pelo Windows
NOMES_RESERVADOS_WINDOWS = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    "COM1",
    "COM2",
    "COM3",
    "COM4",
    "COM5",
    "COM6",
    "COM7",
    "COM8",
    "COM9",
    "LPT1",
    "LPT2",
    "LPT3",
    "LPT4",
    "LPT5",
    "LPT6",
    "LPT7",
    "LPT8",
    "LPT9"
}


# ----------------------------------------------------------------------------
# EXCEÇÕES
# ----------------------------------------------------------------------------

class ExcelValidacaoError(Exception):
    """Erro base para validações da biblioteca Excel."""


class TipoInvalidoError(ExcelValidacaoError, TypeError):
    """Indica que um parâmetro possui tipo incompatível."""


class ValorInvalidoError(ExcelValidacaoError, ValueError):
    """Indica que um parâmetro possui valor inválido."""


class CaminhoInvalidoError(ExcelValidacaoError, ValueError):
    """Indica que um caminho de arquivo ou diretório é inválido."""


class ArquivoExcelNaoEncontradoError(
    ExcelValidacaoError,
    FileNotFoundError
):
    """Indica que o arquivo Excel solicitado não foi encontrado."""


class ExtensaoExcelInvalidaError(ExcelValidacaoError, ValueError):
    """Indica que a extensão do arquivo não é suportada."""


class PlanilhaInvalidaError(ExcelValidacaoError, ValueError):
    """Indica que o nome ou o objeto de planilha é inválido."""


class PlanilhaNaoEncontradaError(ExcelValidacaoError, KeyError):
    """Indica que uma planilha não existe no Workbook."""


class PlanilhaExistenteError(ExcelValidacaoError):
    """Indica que uma planilha já existe no Workbook."""


class ReferenciaExcelInvalidaError(ExcelValidacaoError, ValueError):
    """Indica que uma referência de célula ou intervalo é inválida."""


class DataFrameInvalidoError(ExcelValidacaoError, ValueError):
    """Indica que um DataFrame não atende às regras definidas."""


class ColunaNaoEncontradaError(ExcelValidacaoError, KeyError):
    """Indica que uma coluna obrigatória não existe no DataFrame."""


class CorInvalidaError(ExcelValidacaoError, ValueError):
    """Indica que um código de cor hexadecimal é inválido."""


# ----------------------------------------------------------------------------
# FUNÇÕES DE VALIDAÇÃO DE TIPOS BÁSICOS
# ----------------------------------------------------------------------------

def validar_booleano(
    valor: bool,
    nome: str = "valor"
) -> bool:
    """
    Valida estritamente um parâmetro booleano.

    Args:
        valor (bool):
            Valor que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.

    Returns:
        bool:
            O próprio valor validado.

    Raises:
        TipoInvalidoError:
            Caso o valor não seja booleano.
    """
    # Verifica se o valor informado é booleano
    if not isinstance(valor, bool):
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser booleano."
        )

    # Retorna o valor validado
    return valor


def validar_inteiro(
    valor: int,
    nome: str = "valor",
    minimo: int | None = None,
    maximo: int | None = None
) -> int:
    """
    Valida um número inteiro e seus limites opcionais.

    Args:
        valor (int):
            Número inteiro que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        minimo (int | None):
            Menor valor permitido.
        maximo (int | None):
            Maior valor permitido.

    Returns:
        int:
            O próprio número inteiro validado.

    Raises:
        TipoInvalidoError:
            Caso o valor ou os limites não sejam inteiros.
        ValorInvalidoError:
            Caso o valor esteja fora dos limites permitidos.
    """
    # Impede que booleanos sejam interpretados como inteiros
    if isinstance(valor, bool) or not isinstance(valor, int):
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser inteiro."
        )

    # Valida o limite mínimo quando informado
    if minimo is not None:
        if isinstance(minimo, bool) or not isinstance(minimo, int):
            raise TipoInvalidoError(
                "O parâmetro 'minimo' deve ser inteiro ou None."
            )

        if valor < minimo:
            raise ValorInvalidoError(
                f"O parâmetro '{nome}' deve ser maior ou igual a {minimo}."
            )

    # Valida o limite máximo quando informado
    if maximo is not None:
        if isinstance(maximo, bool) or not isinstance(maximo, int):
            raise TipoInvalidoError(
                "O parâmetro 'maximo' deve ser inteiro ou None."
            )

        if valor > maximo:
            raise ValorInvalidoError(
                f"O parâmetro '{nome}' deve ser menor ou igual a {maximo}."
            )

    # Retorna o número validado
    return valor


def validar_numero(
    valor: int | float,
    nome: str = "valor",
    minimo: int | float | None = None,
    maximo: int | float | None = None
) -> int | float:
    """
    Valida um número inteiro ou decimal e seus limites opcionais.

    Args:
        valor (int | float):
            Número que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        minimo (int | float | None):
            Menor valor permitido.
        maximo (int | float | None):
            Maior valor permitido.

    Returns:
        int | float:
            O próprio valor numérico validado.

    Raises:
        TipoInvalidoError:
            Caso o valor não seja numérico.
        ValorInvalidoError:
            Caso o valor esteja fora dos limites permitidos.
    """
    # Impede que booleanos sejam interpretados como números
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser numérico."
        )

    # Valida o limite mínimo quando informado
    if minimo is not None and valor < minimo:
        raise ValorInvalidoError(
            f"O parâmetro '{nome}' deve ser maior ou igual a {minimo}."
        )

    # Valida o limite máximo quando informado
    if maximo is not None and valor > maximo:
        raise ValorInvalidoError(
            f"O parâmetro '{nome}' deve ser menor ou igual a {maximo}."
        )

    # Retorna o valor validado
    return valor


def validar_texto(
    valor: str,
    nome: str = "valor",
    permitir_vazio: bool = False,
    remover_espacos: bool = True
) -> str:
    """
    Valida e normaliza um parâmetro textual.

    Args:
        valor (str):
            Texto que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        permitir_vazio (bool):
            Define se uma string vazia será aceita.
        remover_espacos (bool):
            Define se os espaços externos serão removidos.

    Returns:
        str:
            Texto validado e normalizado.

    Raises:
        TipoInvalidoError:
            Caso o valor não seja uma string.
        ValorInvalidoError:
            Caso o texto esteja vazio e isso não seja permitido.
    """
    # Valida as opções booleanas
    validar_booleano(
        permitir_vazio,
        "permitir_vazio"
    )
    validar_booleano(
        remover_espacos,
        "remover_espacos"
    )

    # Verifica se o valor informado é uma string
    if not isinstance(valor, str):
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser uma string."
        )

    # Remove os espaços externos quando solicitado
    texto = valor.strip() if remover_espacos else valor

    # Verifica se o texto possui conteúdo
    if not texto and not permitir_vazio:
        raise ValorInvalidoError(
            f"O parâmetro '{nome}' não pode estar vazio."
        )

    # Retorna o texto validado
    return texto


def validar_iteravel(
    valor: Iterable[Any],
    nome: str = "valor",
    permitir_vazio: bool = True,
    permitir_texto: bool = False
) -> list[Any]:
    """
    Valida uma coleção iterável e retorna seus itens em uma lista.

    Args:
        valor (Iterable[Any]):
            Coleção que será validada.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        permitir_vazio (bool):
            Define se uma coleção vazia será aceita.
        permitir_texto (bool):
            Define se strings e bytes serão tratados como iteráveis.

    Returns:
        list[Any]:
            Lista contendo os itens do iterável.

    Raises:
        TipoInvalidoError:
            Caso o valor não seja iterável ou texto não seja permitido.
        ValorInvalidoError:
            Caso a coleção esteja vazia e isso não seja permitido.
    """
    # Valida as opções booleanas
    validar_booleano(
        permitir_vazio,
        "permitir_vazio"
    )
    validar_booleano(
        permitir_texto,
        "permitir_texto"
    )

    # Rejeita textos quando essa opção está desabilitada
    if isinstance(valor, (str, bytes)) and not permitir_texto:
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser um iterável não textual."
        )

    # Verifica se o objeto é iterável
    if not isinstance(valor, Iterable):
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser iterável."
        )

    # Converte a coleção para lista
    itens = list(valor)

    # Verifica se coleções vazias são permitidas
    if not itens and not permitir_vazio:
        raise ValorInvalidoError(
            f"O parâmetro '{nome}' não pode estar vazio."
        )

    # Retorna os itens validados
    return itens


def validar_mapeamento(
    valor: Mapping[Any, Any],
    nome: str = "valor",
    permitir_vazio: bool = True
) -> dict[Any, Any]:
    """
    Valida um objeto de mapeamento e retorna uma cópia em dicionário.

    Args:
        valor (Mapping[Any, Any]):
            Mapeamento que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        permitir_vazio (bool):
            Define se um mapeamento vazio será aceito.

    Returns:
        dict[Any, Any]:
            Cópia do mapeamento validado.

    Raises:
        TipoInvalidoError:
            Caso o valor não seja um mapeamento.
        ValorInvalidoError:
            Caso o mapeamento esteja vazio e isso não seja permitido.
    """
    # Valida a opção de mapeamento vazio
    validar_booleano(
        permitir_vazio,
        "permitir_vazio"
    )

    # Verifica se o objeto é um mapeamento
    if not isinstance(valor, Mapping):
        raise TipoInvalidoError(
            f"O parâmetro '{nome}' deve ser um mapeamento."
        )

    # Cria uma cópia independente do mapeamento
    resultado = dict(valor)

    # Verifica se mapeamentos vazios são permitidos
    if not resultado and not permitir_vazio:
        raise ValorInvalidoError(
            f"O parâmetro '{nome}' não pode estar vazio."
        )

    # Retorna o dicionário validado
    return resultado


# ----------------------------------------------------------------------------
# FUNÇÕES DE VALIDAÇÃO DE CAMINHOS E ARQUIVOS
# ----------------------------------------------------------------------------

def normalizar_caminho(
    caminho: str | Path
) -> Path:
    """
    Normaliza um caminho recebido como string ou Path.

    Args:
        caminho (str | Path):
            Caminho que será normalizado.

    Returns:
        Path:
            Caminho absoluto e normalizado.

    Raises:
        TipoInvalidoError:
            Caso o caminho não seja string ou Path.
        CaminhoInvalidoError:
            Caso uma string vazia seja informada.
    """
    # Verifica se o caminho possui um tipo aceito
    if not isinstance(caminho, (str, Path)):
        raise TipoInvalidoError(
            "O parâmetro 'caminho' deve ser uma string ou Path."
        )

    # Rejeita strings vazias
    if isinstance(caminho, str) and not caminho.strip():
        raise CaminhoInvalidoError(
            "O parâmetro 'caminho' não pode estar vazio."
        )

    # Retorna a representação absoluta do caminho
    return Path(caminho).expanduser().resolve()


def validar_extensao_excel(
    caminho: str | Path,
    extensoes_permitidas: Iterable[str] | None = None
) -> str:
    """
    Valida a extensão de um caminho destinado a arquivos Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo cuja extensão será validada.
        extensoes_permitidas (Iterable[str] | None):
            Coleção opcional de extensões aceitas.

    Returns:
        str:
            Extensão validada em letras minúsculas.

    Raises:
        ExtensaoExcelInvalidaError:
            Caso o caminho não possua uma extensão permitida.
    """
    # Normaliza o caminho recebido
    arquivo = normalizar_caminho(caminho)

    # Define as extensões utilizadas na validação
    if extensoes_permitidas is None:
        extensoes = set(EXTENSOES_EXCEL)
    else:
        itens = validar_iteravel(
            extensoes_permitidas,
            "extensoes_permitidas",
            permitir_vazio=False
        )
        extensoes = {
            validar_texto(
                extensao,
                "extensao"
            ).lower()
            for extensao in itens
        }

    # Obtém a extensão em letras minúsculas
    extensao = arquivo.suffix.lower()

    # Verifica se a extensão está entre as permitidas
    if extensao not in extensoes:
        raise ExtensaoExcelInvalidaError(
            f"A extensão '{extensao}' não é permitida. "
            f"Extensões aceitas: {sorted(extensoes)}."
        )

    # Retorna a extensão validada
    return extensao


def validar_arquivo_excel(
    caminho: str | Path,
    extensoes_permitidas: Iterable[str] | None = None
) -> Path:
    """
    Valida a existência, o tipo e a extensão de um arquivo Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo que será validado.
        extensoes_permitidas (Iterable[str] | None):
            Coleção opcional de extensões aceitas.

    Returns:
        Path:
            Caminho absoluto do arquivo validado.

    Raises:
        ArquivoExcelNaoEncontradoError:
            Caso o arquivo não exista.
        CaminhoInvalidoError:
            Caso o caminho não represente um arquivo.
        ExtensaoExcelInvalidaError:
            Caso a extensão não seja aceita.
    """
    # Normaliza o caminho do arquivo
    arquivo = normalizar_caminho(caminho)

    # Verifica se o caminho existe
    if not arquivo.exists():
        raise ArquivoExcelNaoEncontradoError(
            f"O arquivo Excel não foi encontrado: {arquivo}"
        )

    # Verifica se o caminho representa um arquivo
    if not arquivo.is_file():
        raise CaminhoInvalidoError(
            f"O caminho informado não representa um arquivo: {arquivo}"
        )

    # Valida a extensão do arquivo
    validar_extensao_excel(
        arquivo,
        extensoes_permitidas=extensoes_permitidas
    )

    # Retorna o arquivo validado
    return arquivo


def validar_diretorio(
    caminho: str | Path,
    criar: bool = False
) -> Path:
    """
    Valida um diretório e pode criá-lo quando estiver ausente.

    Args:
        caminho (str | Path):
            Caminho do diretório que será validado.
        criar (bool):
            Define se o diretório será criado quando não existir.

    Returns:
        Path:
            Caminho absoluto do diretório validado.

    Raises:
        FileNotFoundError:
            Caso o diretório não exista e a criação esteja desabilitada.
        CaminhoInvalidoError:
            Caso o caminho exista e não represente um diretório.
    """
    # Normaliza o caminho e valida a opção de criação
    diretorio = normalizar_caminho(caminho)
    validar_booleano(
        criar,
        "criar"
    )

    # Cria o diretório quando solicitado
    if not diretorio.exists() and criar:
        diretorio.mkdir(
            parents=True,
            exist_ok=True
        )

    # Verifica se o diretório existe
    if not diretorio.exists():
        raise FileNotFoundError(
            f"O diretório não foi encontrado: {diretorio}"
        )

    # Verifica se o caminho representa um diretório
    if not diretorio.is_dir():
        raise CaminhoInvalidoError(
            f"O caminho informado não representa um diretório: {diretorio}"
        )

    # Retorna o diretório validado
    return diretorio


# ----------------------------------------------------------------------------
# FUNÇÕES DE VALIDAÇÃO DE WORKBOOK E PLANILHAS
# ----------------------------------------------------------------------------

def validar_workbook(
    workbook: Workbook
) -> Workbook:
    """
    Valida se o objeto informado é um Workbook do openpyxl.

    Args:
        workbook (Workbook):
            Workbook que será validado.

    Returns:
        Workbook:
            O próprio Workbook validado.

    Raises:
        TipoInvalidoError:
            Caso o objeto não seja um Workbook.
    """
    # Verifica se o objeto recebido é um Workbook
    if not isinstance(workbook, Workbook):
        raise TipoInvalidoError(
            "O parâmetro 'workbook' deve ser um Workbook do openpyxl."
        )

    # Retorna o Workbook validado
    return workbook


def validar_planilha(
    planilha: Worksheet
) -> Worksheet:
    """
    Valida se o objeto informado é uma Worksheet do openpyxl.

    Args:
        planilha (Worksheet):
            Planilha que será validada.

    Returns:
        Worksheet:
            A própria planilha validada.

    Raises:
        TipoInvalidoError:
            Caso o objeto não seja uma Worksheet.
    """
    # Verifica se o objeto recebido é uma Worksheet
    if not isinstance(planilha, Worksheet):
        raise TipoInvalidoError(
            "O parâmetro 'planilha' deve ser uma Worksheet do openpyxl."
        )

    # Retorna a planilha validada
    return planilha


def validar_nome_planilha(
    nome: str,
    verificar_tamanho: bool = True
) -> str:
    """
    Valida um nome conforme as regras de planilhas do Excel.

    Args:
        nome (str):
            Nome da planilha que será validado.
        verificar_tamanho (bool):
            Define se o limite de 31 caracteres será validado.

    Returns:
        str:
            Nome validado sem espaços externos.

    Raises:
        PlanilhaInvalidaError:
            Caso o nome esteja vazio, seja longo ou possua caracteres inválidos.
    """
    # Valida a opção de tamanho
    validar_booleano(
        verificar_tamanho,
        "verificar_tamanho"
    )

    # Valida e normaliza o texto
    try:
        nome_validado = validar_texto(
            nome,
            "nome"
        )
    except (TipoInvalidoError, ValorInvalidoError) as erro:
        raise PlanilhaInvalidaError(
            "O nome da planilha é inválido."
        ) from erro

    # Verifica o limite de caracteres do Excel
    if verificar_tamanho and len(nome_validado) > LIMITE_CARACTERES_PLANILHA:
        raise PlanilhaInvalidaError(
            "O nome da planilha não pode ultrapassar 31 caracteres."
        )

    # Localiza os caracteres proibidos presentes no nome
    encontrados = sorted(
        caractere
        for caractere in CARACTERES_INVALIDOS_PLANILHA
        if caractere in nome_validado
    )

    # Rejeita nomes que possuam caracteres proibidos
    if encontrados:
        raise PlanilhaInvalidaError(
            "O nome da planilha possui caracteres inválidos: "
            f"{encontrados}."
        )

    # Rejeita nomes iniciados ou finalizados por apóstrofo
    if nome_validado.startswith("'") or nome_validado.endswith("'"):
        raise PlanilhaInvalidaError(
            "O nome da planilha não pode começar ou terminar com apóstrofo."
        )

    # Retorna o nome validado
    return nome_validado


def validar_planilha_existente(
    workbook: Workbook,
    nome_planilha: str
) -> Worksheet:
    """
    Valida se uma planilha existe no Workbook.

    Args:
        workbook (Workbook):
            Workbook que será consultado.
        nome_planilha (str):
            Nome da planilha procurada.

    Returns:
        Worksheet:
            Planilha encontrada no Workbook.

    Raises:
        PlanilhaNaoEncontradaError:
            Caso a planilha não exista.
    """
    # Valida o Workbook e o nome da planilha
    validar_workbook(workbook)
    nome = validar_nome_planilha(nome_planilha)

    # Verifica se a planilha está disponível
    if nome not in workbook.sheetnames:
        raise PlanilhaNaoEncontradaError(
            f"A planilha '{nome}' não foi encontrada. "
            f"Planilhas disponíveis: {workbook.sheetnames}."
        )

    # Retorna a planilha encontrada
    return workbook[nome]


def validar_planilha_inexistente(
    workbook: Workbook,
    nome_planilha: str
) -> str:
    """
    Valida se um nome de planilha ainda não existe no Workbook.

    Args:
        workbook (Workbook):
            Workbook que será consultado.
        nome_planilha (str):
            Nome que será validado.

    Returns:
        str:
            Nome validado disponível para criação.

    Raises:
        PlanilhaExistenteError:
            Caso o Workbook já possua uma planilha com o nome.
    """
    # Valida o Workbook e o nome da planilha
    validar_workbook(workbook)
    nome = validar_nome_planilha(nome_planilha)

    # Verifica se o nome já está em uso
    if nome in workbook.sheetnames:
        raise PlanilhaExistenteError(
            f"A planilha '{nome}' já existe no Workbook."
        )

    # Retorna o nome disponível
    return nome


# ----------------------------------------------------------------------------
# FUNÇÕES DE VALIDAÇÃO DE CÉLULAS, INTERVALOS E DIMENSÕES
# ----------------------------------------------------------------------------

def validar_referencia_celula(
    referencia: str,
    planilha: Worksheet | None = None,
    permitir_mesclada: bool = False
) -> str:
    """
    Valida e normaliza uma referência de célula única.

    Args:
        referencia (str):
            Referência da célula no padrão do Excel.
        planilha (Worksheet | None):
            Planilha opcional utilizada para validar células mescladas.
        permitir_mesclada (bool):
            Define se células mescladas secundárias serão aceitas.

    Returns:
        str:
            Referência validada em letras maiúsculas.

    Raises:
        ReferenciaExcelInvalidaError:
            Caso a referência não represente uma célula válida.
    """
    # Valida a opção de célula mesclada
    validar_booleano(
        permitir_mesclada,
        "permitir_mesclada"
    )

    # Valida e normaliza o texto
    coordenada = validar_texto(
        referencia,
        "referencia"
    ).upper()

    # Obtém os limites da referência
    try:
        limites = range_boundaries(coordenada)
    except (TypeError, ValueError) as erro:
        raise ReferenciaExcelInvalidaError(
            f"A referência de célula é inválida: '{referencia}'."
        ) from erro

    # Rejeita referências parciais ou com índice zero
    if any(limite is None for limite in limites) or min(limites) < 1:
        raise ReferenciaExcelInvalidaError(
            f"A referência de célula é inválida: '{referencia}'."
        )

    # Extrai os limites da referência
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites

    # Exige uma única célula
    if coluna_inicial != coluna_final or linha_inicial != linha_final:
        raise ReferenciaExcelInvalidaError(
            "A referência deve representar uma única célula."
        )

    # Verifica os limites estruturais do Excel
    if linha_inicial > LIMITE_LINHAS_EXCEL or coluna_inicial > LIMITE_COLUNAS_EXCEL:
        raise ReferenciaExcelInvalidaError(
            "A referência ultrapassa os limites estruturais do Excel."
        )

    # Valida a célula em uma planilha quando ela foi informada
    if planilha is not None:
        validar_planilha(planilha)
        celula = planilha[coordenada]

        # Rejeita células mescladas secundárias quando necessário
        if isinstance(celula, MergedCell) and not permitir_mesclada:
            raise ReferenciaExcelInvalidaError(
                f"A célula '{coordenada}' pertence a um intervalo mesclado."
            )

    # Retorna a referência normalizada
    return coordenada


def validar_referencia_intervalo(
    referencia: str
) -> str:
    """
    Valida e normaliza uma referência de célula ou intervalo.

    Args:
        referencia (str):
            Referência no padrão do Excel, como A1 ou A1:D10.

    Returns:
        str:
            Referência validada em letras maiúsculas.

    Raises:
        ReferenciaExcelInvalidaError:
            Caso a referência seja incompleta ou ultrapasse os limites.
    """
    # Valida e normaliza o texto
    intervalo = validar_texto(
        referencia,
        "referencia"
    ).upper()

    # Obtém os limites da referência
    try:
        limites = range_boundaries(intervalo)
    except (TypeError, ValueError) as erro:
        raise ReferenciaExcelInvalidaError(
            f"A referência de intervalo é inválida: '{referencia}'."
        ) from erro

    # Rejeita referências parciais ou índices menores que um
    if any(limite is None for limite in limites) or min(limites) < 1:
        raise ReferenciaExcelInvalidaError(
            f"A referência de intervalo é inválida: '{referencia}'."
        )

    # Extrai os limites máximos
    _, _, coluna_final, linha_final = limites

    # Verifica os limites estruturais do Excel
    if linha_final > LIMITE_LINHAS_EXCEL or coluna_final > LIMITE_COLUNAS_EXCEL:
        raise ReferenciaExcelInvalidaError(
            "A referência ultrapassa os limites estruturais do Excel."
        )

    # Retorna a referência validada
    return intervalo


def validar_linha(
    linha: int
) -> int:
    """
    Valida um índice de linha conforme os limites do Excel.

    Args:
        linha (int):
            Número da linha iniciado em um.

    Returns:
        int:
            Índice da linha validado.

    Raises:
        ValorInvalidoError:
            Caso a linha esteja fora dos limites do Excel.
    """
    # Reutiliza a validação centralizada de inteiros
    return validar_inteiro(
        linha,
        "linha",
        minimo=1,
        maximo=LIMITE_LINHAS_EXCEL
    )


def validar_coluna(
    coluna: str | int,
    retornar_indice: bool = False
) -> str | int:
    """
    Valida uma coluna informada por letra ou índice numérico.

    Args:
        coluna (str | int):
            Letra ou índice da coluna.
        retornar_indice (bool):
            Define se o retorno será o índice numérico da coluna.

    Returns:
        str | int:
            Letra normalizada ou índice numérico da coluna.

    Raises:
        TipoInvalidoError:
            Caso a coluna não seja string ou inteiro.
        ValorInvalidoError:
            Caso a coluna ultrapasse os limites do Excel.
    """
    # Valida a opção de retorno
    validar_booleano(
        retornar_indice,
        "retornar_indice"
    )

    # Trata a coluna informada por índice
    if isinstance(coluna, int) and not isinstance(coluna, bool):
        indice = validar_inteiro(
            coluna,
            "coluna",
            minimo=1,
            maximo=LIMITE_COLUNAS_EXCEL
        )
        letra = get_column_letter(indice)

    # Trata a coluna informada por letra
    elif isinstance(coluna, str):
        letra = validar_texto(
            coluna,
            "coluna"
        ).upper()

        # Rejeita referências com números ou símbolos
        if not letra.isalpha():
            raise ValorInvalidoError(
                "A coluna textual deve possuir somente letras."
            )

        # Converte a letra para índice
        try:
            indice = column_index_from_string(letra)
        except ValueError as erro:
            raise ValorInvalidoError(
                f"A coluna informada é inválida: '{coluna}'."
            ) from erro

        # Verifica o limite máximo do Excel
        if indice > LIMITE_COLUNAS_EXCEL:
            raise ValorInvalidoError(
                f"A coluna não pode ultrapassar {ULTIMA_COLUNA_EXCEL}."
            )
    else:
        raise TipoInvalidoError(
            "O parâmetro 'coluna' deve ser uma string ou inteiro."
        )

    # Retorna o formato solicitado
    return indice if retornar_indice else letra


# ----------------------------------------------------------------------------
# FUNÇÕES DE VALIDAÇÃO DE DATAFRAMES E FORMATOS
# ----------------------------------------------------------------------------

def validar_dataframe(
    dataframe: pd.DataFrame,
    permitir_vazio: bool = True,
    exigir_colunas: bool = True,
    permitir_colunas_duplicadas: bool = False
) -> pd.DataFrame:
    """
    Valida um DataFrame utilizado nas operações da biblioteca Excel.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será validado.
        permitir_vazio (bool):
            Define se DataFrames sem registros serão aceitos.
        exigir_colunas (bool):
            Define se ao menos uma coluna deve existir.
        permitir_colunas_duplicadas (bool):
            Define se nomes de colunas duplicados serão aceitos.

    Returns:
        pd.DataFrame:
            O próprio DataFrame validado.

    Raises:
        DataFrameInvalidoError:
            Caso o objeto ou sua estrutura seja inválida.
    """
    # Valida as opções booleanas
    validar_booleano(
        permitir_vazio,
        "permitir_vazio"
    )
    validar_booleano(
        exigir_colunas,
        "exigir_colunas"
    )
    validar_booleano(
        permitir_colunas_duplicadas,
        "permitir_colunas_duplicadas"
    )

    # Verifica se o objeto é um DataFrame
    if not isinstance(dataframe, pd.DataFrame):
        raise DataFrameInvalidoError(
            "O objeto fornecido não é um DataFrame do pandas."
        )

    # Verifica se existem colunas quando elas são obrigatórias
    if exigir_colunas and len(dataframe.columns) == 0:
        raise DataFrameInvalidoError(
            "O DataFrame deve possuir ao menos uma coluna."
        )

    # Verifica se existem registros quando eles são obrigatórios
    if not permitir_vazio and dataframe.empty:
        raise DataFrameInvalidoError(
            "O DataFrame não pode estar vazio."
        )

    # Identifica nomes duplicados
    duplicadas = (
        dataframe.columns[
            dataframe.columns.duplicated()
        ]
        .astype(str)
        .unique()
        .tolist()
    )

    # Rejeita duplicidades quando necessário
    if duplicadas and not permitir_colunas_duplicadas:
        raise DataFrameInvalidoError(
            f"O DataFrame possui colunas duplicadas: {duplicadas}."
        )

    # Retorna o DataFrame validado
    return dataframe


def validar_colunas_dataframe(
    dataframe: pd.DataFrame,
    colunas_obrigatorias: Iterable[str],
    considerar_maiusculas: bool = True
) -> list[str]:
    """
    Valida a existência de colunas obrigatórias em um DataFrame.

    Args:
        dataframe (pd.DataFrame):
            DataFrame que será consultado.
        colunas_obrigatorias (Iterable[str]):
            Coleção com os nomes das colunas esperadas.
        considerar_maiusculas (bool):
            Define se a comparação diferencia maiúsculas e minúsculas.

    Returns:
        list[str]:
            Lista com os nomes validados na ordem informada.

    Raises:
        ColunaNaoEncontradaError:
            Caso alguma coluna obrigatória esteja ausente.
    """
    # Valida o DataFrame e a opção de comparação
    validar_dataframe(dataframe)
    validar_booleano(
        considerar_maiusculas,
        "considerar_maiusculas"
    )

    # Valida a coleção de colunas obrigatórias
    itens = validar_iteravel(
        colunas_obrigatorias,
        "colunas_obrigatorias",
        permitir_vazio=False
    )

    # Valida os nomes das colunas
    colunas = [
        validar_texto(
            coluna,
            "coluna"
        )
        for coluna in itens
    ]

    # Prepara os nomes disponíveis conforme a regra de comparação
    if considerar_maiusculas:
        disponiveis = {
            str(coluna)
            for coluna in dataframe.columns
        }
        ausentes = [
            coluna
            for coluna in colunas
            if coluna not in disponiveis
        ]
    else:
        disponiveis = {
            str(coluna).casefold()
            for coluna in dataframe.columns
        }
        ausentes = [
            coluna
            for coluna in colunas
            if coluna.casefold() not in disponiveis
        ]

    # Rejeita a ausência de colunas obrigatórias
    if ausentes:
        raise ColunaNaoEncontradaError(
            f"As colunas obrigatórias não foram encontradas: {ausentes}."
        )

    # Retorna os nomes validados
    return colunas


def validar_cor_hexadecimal(
    cor: str,
    retornar_argb: bool = True
) -> str:
    """
    Valida uma cor hexadecimal RGB ou ARGB.

    Args:
        cor (str):
            Código de cor com seis ou oito caracteres hexadecimais.
        retornar_argb (bool):
            Define se cores RGB receberão o prefixo alfa FF.

    Returns:
        str:
            Código hexadecimal validado em letras maiúsculas.

    Raises:
        CorInvalidaError:
            Caso o código possua tamanho ou caracteres inválidos.
    """
    # Valida a opção de retorno
    validar_booleano(
        retornar_argb,
        "retornar_argb"
    )

    # Valida o texto e remove o prefixo opcional
    codigo = validar_texto(
        cor,
        "cor"
    ).lstrip("#").upper()

    # Verifica se todos os caracteres são hexadecimais
    if any(
        caractere not in "0123456789ABCDEF"
        for caractere in codigo
    ):
        raise CorInvalidaError(
            f"O código de cor é inválido: '{cor}'."
        )

    # Verifica o tamanho permitido
    if len(codigo) not in {6, 8}:
        raise CorInvalidaError(
            "A cor deve possuir seis caracteres RGB ou oito caracteres ARGB."
        )

    # Adiciona o canal alfa para cores RGB quando solicitado
    if len(codigo) == 6 and retornar_argb:
        codigo = f"FF{codigo}"

    # Retorna a cor validada
    return codigo


def validar_formato_numero(
    formato: str
) -> str:
    """
    Valida um código textual de formato numérico do Excel.

    Args:
        formato (str):
            Código de formatação numérica.

    Returns:
        str:
            Formato validado sem espaços externos.

    Raises:
        ValorInvalidoError:
            Caso o formato contenha caractere nulo.
    """
    # Valida o texto do formato
    formato_validado = validar_texto(
        formato,
        "formato"
    )

    # Rejeita caracteres nulos
    if "\x00" in formato_validado:
        raise ValorInvalidoError(
            "O formato numérico possui caracteres inválidos."
        )

    # Retorna o formato validado
    return formato_validado


def validar_nome_tabela(
    nome: str
) -> str:
    """
    Valida um nome destinado a uma tabela estruturada do Excel.

    Args:
        nome (str):
            Nome interno da tabela estruturada.

    Returns:
        str:
            Nome validado da tabela.

    Raises:
        ValorInvalidoError:
            Caso o nome possua espaços, símbolos ou inicie com número.
    """
    # Valida o texto do nome
    nome_validado = validar_texto(
        nome,
        "nome"
    )

    # Rejeita nomes iniciados por número
    if nome_validado[0].isdigit():
        raise ValorInvalidoError(
            "O nome da tabela não pode começar com um número."
        )

    # Rejeita caracteres diferentes de letras, números e sublinhado
    if not all(
        caractere.isalnum() or caractere == "_"
        for caractere in nome_validado
    ):
        raise ValorInvalidoError(
            "O nome da tabela deve possuir somente letras, números e sublinhado."
        )

    # Retorna o nome validado
    return nome_validado


# ----------------------------------------------------------------------------
# EXPORTAÇÕES PÚBLICAS
# ----------------------------------------------------------------------------

__all__ = [
    "EXTENSOES_EXCEL",
    "EXTENSOES_OPENPYXL",
    "LIMITE_LINHAS_EXCEL",
    "LIMITE_COLUNAS_EXCEL",
    "LIMITE_CARACTERES_PLANILHA",
    "LIMITE_CARACTERES_CELULA",
    "ULTIMA_COLUNA_EXCEL",
    "CARACTERES_INVALIDOS_PLANILHA",
    "NOMES_RESERVADOS_WINDOWS",
    "ExcelValidacaoError",
    "TipoInvalidoError",
    "ValorInvalidoError",
    "CaminhoInvalidoError",
    "ArquivoExcelNaoEncontradoError",
    "ExtensaoExcelInvalidaError",
    "PlanilhaInvalidaError",
    "PlanilhaNaoEncontradaError",
    "PlanilhaExistenteError",
    "ReferenciaExcelInvalidaError",
    "DataFrameInvalidoError",
    "ColunaNaoEncontradaError",
    "CorInvalidaError",
    "validar_booleano",
    "validar_inteiro",
    "validar_numero",
    "validar_texto",
    "validar_iteravel",
    "validar_mapeamento",
    "normalizar_caminho",
    "validar_extensao_excel",
    "validar_arquivo_excel",
    "validar_diretorio",
    "validar_workbook",
    "validar_planilha",
    "validar_nome_planilha",
    "validar_planilha_existente",
    "validar_planilha_inexistente",
    "validar_referencia_celula",
    "validar_referencia_intervalo",
    "validar_linha",
    "validar_coluna",
    "validar_dataframe",
    "validar_colunas_dataframe",
    "validar_cor_hexadecimal",
    "validar_formato_numero",
    "validar_nome_tabela"
]
