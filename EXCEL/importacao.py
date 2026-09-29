"""
============================================================
Módulo: EXCEL / importacao.py
Autor: Tiago Eneas Antunes (Z701038)
Data de Criação: 28/09/2026
Última Alteração: 28/09/2026
Versão: 1.0.0
Descrição:
    Biblioteca de funções auxiliares para importar arquivos Excel como
    DataFrames, registros e estruturas de dados. O módulo utiliza pandas
    e openpyxl para listar planilhas, validar arquivos, ler intervalos,
    importar múltiplas abas e consultar metadados.
Funções disponíveis:
    - validar_arquivo_excel()
    - listar_planilhas()
    - planilha_existe()
    - obter_metadados_arquivo()
    - importar_dataframe()
    - importar_multiplas_planilhas()
    - importar_todas_planilhas()
    - importar_intervalo()
    - importar_colunas()
    - importar_registros()
    - importar_cabecalhos()
    - localizar_linha_cabecalho()
    - importar_tabela_excel()
Dependências:
    - pandas
    - openpyxl
    - pathlib
Histórico:
    v1.0.0 - 28/09/2026
        - Criação inicial do módulo.
        - Inclusão da validação de arquivos Excel.
        - Inclusão da importação de uma ou várias planilhas.
        - Inclusão da leitura por intervalo, colunas e tabela estruturada.
        - Inclusão da consulta de cabeçalhos e metadados.
============================================================
"""

# ----------------------------------------------------------------------------
# IMPORTAÇÕES
# ----------------------------------------------------------------------------

# Importa Path para manipulação segura dos caminhos dos arquivos
from pathlib import Path

# Importa Any e Iterable para anotações de tipos flexíveis
from typing import Any, Iterable

# Importa pandas para leitura e organização dos dados em DataFrames
import pandas as pd

# Importa load_workbook para consultas estruturais aos arquivos Excel
from openpyxl import load_workbook

# Importa range_boundaries para interpretar intervalos do Excel
from openpyxl.utils import range_boundaries


# ----------------------------------------------------------------------------
# CONSTANTES
# ----------------------------------------------------------------------------

# Define as extensões reconhecidas como arquivos Excel
EXTENSOES_EXCEL = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
    ".xls"
}

# Define as extensões lidas pelo mecanismo openpyxl
EXTENSOES_OPENPYXL = {
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm"
}

# Define a extensão legada que exige mecanismo específico do pandas
EXTENSAO_EXCEL_LEGADA = ".xls"

# Define o mecanismo padrão para os formatos modernos
ENGINE_EXCEL_PADRAO = "openpyxl"

# Define o mecanismo padrão para o formato legado
ENGINE_EXCEL_LEGADO = "xlrd"

# Define a linha padrão utilizada como cabeçalho pelo pandas
LINHA_CABECALHO_PADRAO = 0


# ----------------------------------------------------------------------------
# EXCEÇÕES
# ----------------------------------------------------------------------------

class ExcelImportacaoError(Exception):
    """Erro base para operações de importação de arquivos Excel."""


class ArquivoExcelInvalidoError(ExcelImportacaoError, ValueError):
    """Indica que o caminho, a extensão ou o conteúdo Excel é inválido."""


class ArquivoExcelNaoEncontradoError(
    ExcelImportacaoError,
    FileNotFoundError
):
    """Indica que o arquivo Excel solicitado não foi encontrado."""


class PlanilhaNaoEncontradaError(ExcelImportacaoError, KeyError):
    """Indica que a planilha solicitada não existe no arquivo Excel."""


class IntervaloImportacaoInvalidoError(ExcelImportacaoError, ValueError):
    """Indica que o intervalo solicitado para importação é inválido."""


class TabelaExcelNaoEncontradaError(ExcelImportacaoError, KeyError):
    """Indica que a tabela estruturada solicitada não foi encontrada."""


class ImportacaoDataFrameError(ExcelImportacaoError):
    """Indica que os dados não puderam ser importados como DataFrame."""


# ----------------------------------------------------------------------------
# FUNÇÕES INTERNAS DE VALIDAÇÃO
# ----------------------------------------------------------------------------

def _validar_booleano(
    valor: bool,
    nome: str
) -> bool:
    """
    Valida estritamente um parâmetro booleano.

    Args:
        valor (bool):
            Valor que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.

    Returns:
        bool:
            O próprio valor validado.

    Raises:
        TypeError:
            Caso o valor não seja booleano.
    """
    # Verifica se o valor informado é booleano
    if not isinstance(valor, bool):
        raise TypeError(
            f"O parâmetro '{nome}' deve ser booleano."
        )

    # Retorna o valor validado
    return valor


def _validar_inteiro(
    valor: int,
    nome: str,
    minimo: int = 0
) -> int:
    """
    Valida um número inteiro e seu limite mínimo.

    Args:
        valor (int):
            Número inteiro que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.
        minimo (int):
            Menor valor permitido para o parâmetro.

    Returns:
        int:
            O próprio número inteiro validado.

    Raises:
        TypeError:
            Caso o valor não seja inteiro.
        ValueError:
            Caso o valor seja menor que o limite mínimo.
    """
    # Impede que valores booleanos sejam interpretados como inteiros
    if isinstance(valor, bool) or not isinstance(valor, int):
        raise TypeError(
            f"O parâmetro '{nome}' deve ser inteiro."
        )

    # Verifica se o valor respeita o limite mínimo
    if valor < minimo:
        raise ValueError(
            f"O parâmetro '{nome}' deve ser maior ou igual a {minimo}."
        )

    # Retorna o número validado
    return valor


def _validar_texto(
    valor: str,
    nome: str
) -> str:
    """
    Valida e normaliza um parâmetro textual obrigatório.

    Args:
        valor (str):
            Texto que será validado.
        nome (str):
            Nome do parâmetro utilizado na mensagem de erro.

    Returns:
        str:
            Texto sem espaços externos.

    Raises:
        TypeError:
            Caso o valor não seja uma string.
        ValueError:
            Caso o texto esteja vazio.
    """
    # Verifica se o valor informado é uma string
    if not isinstance(valor, str):
        raise TypeError(
            f"O parâmetro '{nome}' deve ser uma string."
        )

    # Remove os espaços externos
    texto = valor.strip()

    # Verifica se o texto possui conteúdo
    if not texto:
        raise ValueError(
            f"O parâmetro '{nome}' não pode estar vazio."
        )

    # Retorna o texto normalizado
    return texto


def _normalizar_caminho(
    caminho: str | Path
) -> Path:
    """
    Normaliza um caminho recebido como string ou Path.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.

    Returns:
        Path:
            Caminho absoluto e normalizado.

    Raises:
        TypeError:
            Caso o caminho não seja string ou Path.
        ValueError:
            Caso uma string vazia seja informada.
    """
    # Verifica se o caminho possui um tipo aceito
    if not isinstance(caminho, (str, Path)):
        raise TypeError(
            "O parâmetro 'caminho' deve ser uma string ou Path."
        )

    # Rejeita strings vazias
    if isinstance(caminho, str) and not caminho.strip():
        raise ValueError(
            "O parâmetro 'caminho' não pode estar vazio."
        )

    # Retorna o caminho absoluto normalizado
    return Path(caminho).expanduser().resolve()


def _obter_engine(
    caminho: Path
) -> str:
    """
    Retorna o mecanismo adequado para a extensão do arquivo.

    Args:
        caminho (Path):
            Caminho do arquivo Excel validado.

    Returns:
        str:
            Nome do mecanismo utilizado pelo pandas.

    Raises:
        ArquivoExcelInvalidoError:
            Caso a extensão não seja suportada.
    """
    # Obtém a extensão em letras minúsculas
    extensao = caminho.suffix.lower()

    # Retorna o mecanismo dos arquivos modernos
    if extensao in EXTENSOES_OPENPYXL:
        return ENGINE_EXCEL_PADRAO

    # Retorna o mecanismo do formato legado
    if extensao == EXTENSAO_EXCEL_LEGADA:
        return ENGINE_EXCEL_LEGADO

    # Interrompe a execução para extensões desconhecidas
    raise ArquivoExcelInvalidoError(
        f"A extensão do arquivo não é suportada: '{extensao}'."
    )


def _normalizar_intervalo(
    intervalo: str
) -> tuple[str, tuple[int, int, int, int]]:
    """
    Valida um intervalo e retorna sua referência e limites numéricos.

    Args:
        intervalo (str):
            Intervalo no padrão do Excel, como A1:D10.

    Returns:
        tuple[str, tuple[int, int, int, int]]:
            Referência normalizada e limites do intervalo.

    Raises:
        IntervaloImportacaoInvalidoError:
            Caso a referência seja incompleta ou inválida.
    """
    # Valida o texto informado
    referencia = _validar_texto(
        intervalo,
        "intervalo"
    ).upper()

    # Tenta interpretar os limites da referência
    try:
        limites = range_boundaries(referencia)
    except (TypeError, ValueError) as erro:
        raise IntervaloImportacaoInvalidoError(
            f"O intervalo informado é inválido: '{intervalo}'."
        ) from erro

    # Rejeita referências sem linha ou coluna
    if any(limite is None for limite in limites):
        raise IntervaloImportacaoInvalidoError(
            f"O intervalo informado é inválido: '{intervalo}'."
        )

    # Rejeita índices menores que um
    if min(limites) < 1:
        raise IntervaloImportacaoInvalidoError(
            f"O intervalo informado é inválido: '{intervalo}'."
        )

    # Retorna a referência e seus limites
    return referencia, limites


def _validar_planilha_existente(
    caminho: Path,
    nome_planilha: str
) -> str:
    """
    Valida se uma planilha existe no arquivo informado.

    Args:
        caminho (Path):
            Caminho do arquivo Excel validado.
        nome_planilha (str):
            Nome da planilha procurada.

    Returns:
        str:
            Nome validado da planilha.

    Raises:
        PlanilhaNaoEncontradaError:
            Caso a planilha não exista no arquivo.
    """
    # Valida o texto do nome da planilha
    nome = _validar_texto(
        nome_planilha,
        "nome_planilha"
    )

    # Obtém os nomes das planilhas disponíveis
    planilhas = listar_planilhas(caminho)

    # Verifica se o nome existe exatamente no arquivo
    if nome not in planilhas:
        raise PlanilhaNaoEncontradaError(
            f"A planilha '{nome}' não foi encontrada. "
            f"Planilhas disponíveis: {planilhas}."
        )

    # Retorna o nome validado
    return nome


# ----------------------------------------------------------------------------
# FUNÇÕES DE ARQUIVOS E PLANILHAS
# ----------------------------------------------------------------------------

def validar_arquivo_excel(
    caminho: str | Path,
    permitir_xls: bool = True
) -> Path:
    """
    Valida a existência e a extensão de um arquivo Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo que será validado.
        permitir_xls (bool):
            Define se o formato legado .xls será aceito.

    Returns:
        Path:
            Caminho absoluto do arquivo validado.

    Raises:
        ArquivoExcelNaoEncontradoError:
            Caso o arquivo não exista.
        ArquivoExcelInvalidoError:
            Caso o caminho não seja um arquivo ou a extensão seja inválida.
    """
    # Normaliza o caminho recebido
    arquivo = _normalizar_caminho(caminho)

    # Valida a opção do formato legado
    _validar_booleano(
        permitir_xls,
        "permitir_xls"
    )

    # Verifica se o caminho existe
    if not arquivo.exists():
        raise ArquivoExcelNaoEncontradoError(
            f"O arquivo Excel não foi encontrado: {arquivo}"
        )

    # Verifica se o caminho representa um arquivo
    if not arquivo.is_file():
        raise ArquivoExcelInvalidoError(
            f"O caminho informado não representa um arquivo: {arquivo}"
        )

    # Obtém a extensão do arquivo
    extensao = arquivo.suffix.lower()

    # Verifica se a extensão é reconhecida
    if extensao not in EXTENSOES_EXCEL:
        raise ArquivoExcelInvalidoError(
            f"A extensão do arquivo não é suportada: '{extensao}'."
        )

    # Verifica se o formato legado foi permitido
    if extensao == EXTENSAO_EXCEL_LEGADA and not permitir_xls:
        raise ArquivoExcelInvalidoError(
            "O formato legado .xls não está permitido nesta operação."
        )

    # Retorna o arquivo validado
    return arquivo


def listar_planilhas(
    caminho: str | Path
) -> list[str]:
    """
    Lista os nomes das planilhas existentes em um arquivo Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.

    Returns:
        list[str]:
            Nomes das planilhas na ordem existente no arquivo.

    Raises:
        ImportacaoDataFrameError:
            Caso o arquivo não possa ser aberto.
    """
    # Valida o arquivo recebido
    arquivo = validar_arquivo_excel(caminho)

    # Obtém o mecanismo adequado para o arquivo
    engine = _obter_engine(arquivo)

    # Abre o arquivo apenas para consultar os nomes das planilhas
    try:
        with pd.ExcelFile(
            arquivo,
            engine=engine
        ) as excel:
            planilhas = list(excel.sheet_names)
    except Exception as erro:
        raise ImportacaoDataFrameError(
            f"Não foi possível listar as planilhas do arquivo: {arquivo}"
        ) from erro

    # Retorna os nomes encontrados
    return planilhas


def planilha_existe(
    caminho: str | Path,
    nome_planilha: str,
    considerar_maiusculas: bool = True
) -> bool:
    """
    Verifica se uma planilha existe no arquivo Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str):
            Nome da planilha procurada.
        considerar_maiusculas (bool):
            Define se a comparação diferencia maiúsculas e minúsculas.

    Returns:
        bool:
            True quando a planilha existe; caso contrário, False.

    Raises:
        TypeError:
            Caso algum parâmetro possua tipo incompatível.
        ValueError:
            Caso o nome esteja vazio.
    """
    # Valida o nome e a opção de comparação
    nome = _validar_texto(
        nome_planilha,
        "nome_planilha"
    )
    _validar_booleano(
        considerar_maiusculas,
        "considerar_maiusculas"
    )

    # Obtém as planilhas existentes
    planilhas = listar_planilhas(caminho)

    # Realiza a comparação exata quando solicitado
    if considerar_maiusculas:
        return nome in planilhas

    # Compara os nomes sem diferenciar maiúsculas e minúsculas
    return nome.casefold() in {
        planilha.casefold()
        for planilha in planilhas
    }


def obter_metadados_arquivo(
    caminho: str | Path
) -> dict[str, Any]:
    """
    Retorna metadados básicos de um arquivo Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.

    Returns:
        dict[str, Any]:
            Dicionário com caminho, extensão, tamanho e planilhas.

    Raises:
        ArquivoExcelInvalidoError:
            Caso o arquivo não seja válido.
    """
    # Valida o arquivo e obtém as planilhas
    arquivo = validar_arquivo_excel(caminho)
    planilhas = listar_planilhas(arquivo)

    # Cria a estrutura de metadados
    metadados = {
        "nome": arquivo.name,
        "caminho": arquivo,
        "diretorio": arquivo.parent,
        "extensao": arquivo.suffix.lower(),
        "tamanho_bytes": arquivo.stat().st_size,
        "planilhas": planilhas,
        "quantidade_planilhas": len(planilhas),
        "possui_macro": arquivo.suffix.lower() in {".xlsm", ".xltm"},
        "engine": _obter_engine(arquivo)
    }

    # Retorna os dados coletados
    return metadados


# ----------------------------------------------------------------------------
# FUNÇÕES DE IMPORTAÇÃO DE DATAFRAMES
# ----------------------------------------------------------------------------

def importar_dataframe(
    caminho: str | Path,
    nome_planilha: str | int = 0,
    linha_cabecalho: int | None = LINHA_CABECALHO_PADRAO,
    colunas: str | list[str] | None = None,
    quantidade_linhas: int | None = None,
    pular_linhas: int | list[int] | None = None,
    tipos: dict[str, Any] | None = None,
    manter_colunas_vazias: bool = True
) -> pd.DataFrame:
    """
    Importa uma planilha do Excel como DataFrame.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str | int):
            Nome ou índice da planilha que será importada.
        linha_cabecalho (int | None):
            Linha usada como cabeçalho pelo pandas, iniciada em zero.
        colunas (str | list[str] | None):
            Seleção opcional de colunas aceita pelo parâmetro usecols.
        quantidade_linhas (int | None):
            Quantidade máxima de linhas importadas.
        pular_linhas (int | list[int] | None):
            Linhas que serão ignoradas durante a leitura.
        tipos (dict[str, Any] | None):
            Mapeamento opcional de tipos das colunas.
        manter_colunas_vazias (bool):
            Define se colunas totalmente vazias serão preservadas.

    Returns:
        pd.DataFrame:
            DataFrame contendo os dados importados.

    Raises:
        PlanilhaNaoEncontradaError:
            Caso a planilha textual não exista.
        ImportacaoDataFrameError:
            Caso o pandas não consiga realizar a leitura.
    """
    # Valida o arquivo recebido
    arquivo = validar_arquivo_excel(caminho)

    # Valida a referência da planilha
    if isinstance(nome_planilha, str):
        planilha = _validar_planilha_existente(
            arquivo,
            nome_planilha
        )
    elif isinstance(nome_planilha, int) and not isinstance(nome_planilha, bool):
        planilha = _validar_inteiro(
            nome_planilha,
            "nome_planilha",
            minimo=0
        )

        # Valida o índice contra a quantidade de planilhas
        planilhas = listar_planilhas(arquivo)
        if planilha >= len(planilhas):
            raise PlanilhaNaoEncontradaError(
                f"O índice da planilha é inválido: {planilha}."
            )
    else:
        raise TypeError(
            "O parâmetro 'nome_planilha' deve ser string ou inteiro."
        )

    # Valida os parâmetros numéricos opcionais
    if linha_cabecalho is not None:
        _validar_inteiro(
            linha_cabecalho,
            "linha_cabecalho",
            minimo=0
        )

    if quantidade_linhas is not None:
        _validar_inteiro(
            quantidade_linhas,
            "quantidade_linhas",
            minimo=1
        )

    # Valida o mapeamento opcional de tipos
    if tipos is not None and not isinstance(tipos, dict):
        raise TypeError(
            "O parâmetro 'tipos' deve ser um dicionário ou None."
        )

    # Valida a opção das colunas vazias
    _validar_booleano(
        manter_colunas_vazias,
        "manter_colunas_vazias"
    )

    # Executa a leitura através do pandas
    try:
        dataframe = pd.read_excel(
            arquivo,
            sheet_name=planilha,
            header=linha_cabecalho,
            usecols=colunas,
            nrows=quantidade_linhas,
            skiprows=pular_linhas,
            dtype=tipos,
            engine=_obter_engine(arquivo)
        )
    except Exception as erro:
        raise ImportacaoDataFrameError(
            f"Não foi possível importar a planilha '{planilha}' "
            f"do arquivo: {arquivo}"
        ) from erro

    # Remove colunas totalmente vazias quando solicitado
    if not manter_colunas_vazias:
        dataframe = dataframe.dropna(
            axis=1,
            how="all"
        )

    # Retorna uma cópia independente do DataFrame importado
    return dataframe.copy(deep=True)


def importar_multiplas_planilhas(
    caminho: str | Path,
    planilhas: Iterable[str],
    linha_cabecalho: int | None = LINHA_CABECALHO_PADRAO
) -> dict[str, pd.DataFrame]:
    """
    Importa planilhas específicas como um dicionário de DataFrames.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        planilhas (Iterable[str]):
            Coleção com os nomes das planilhas importadas.
        linha_cabecalho (int | None):
            Linha usada como cabeçalho pelo pandas.

    Returns:
        dict[str, pd.DataFrame]:
            Dicionário associando cada planilha ao seu DataFrame.

    Raises:
        TypeError:
            Caso planilhas não seja um iterável textual válido.
        ValueError:
            Caso a coleção esteja vazia ou possua nomes duplicados.
    """
    # Rejeita uma string isolada para evitar uma planilha por caractere
    if isinstance(planilhas, (str, bytes)) or not isinstance(
        planilhas,
        Iterable
    ):
        raise TypeError(
            "O parâmetro 'planilhas' deve ser um iterável não textual."
        )

    # Converte a coleção para lista
    nomes = list(planilhas)

    # Verifica se ao menos uma planilha foi informada
    if not nomes:
        raise ValueError(
            "O parâmetro 'planilhas' não pode estar vazio."
        )

    # Valida todos os nomes
    nomes_validados = [
        _validar_texto(nome, "nome_planilha")
        for nome in nomes
    ]

    # Verifica duplicidades na lista
    if len(nomes_validados) != len(set(nomes_validados)):
        raise ValueError(
            "A coleção de planilhas não pode possuir nomes duplicados."
        )

    # Valida o arquivo uma única vez
    arquivo = validar_arquivo_excel(caminho)

    # Cria o dicionário de resultados
    resultado: dict[str, pd.DataFrame] = {}

    # Importa cada planilha solicitada
    for nome in nomes_validados:
        resultado[nome] = importar_dataframe(
            arquivo,
            nome_planilha=nome,
            linha_cabecalho=linha_cabecalho
        )

    # Retorna os DataFrames importados
    return resultado


def importar_todas_planilhas(
    caminho: str | Path,
    linha_cabecalho: int | None = LINHA_CABECALHO_PADRAO
) -> dict[str, pd.DataFrame]:
    """
    Importa todas as planilhas de um arquivo Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        linha_cabecalho (int | None):
            Linha usada como cabeçalho pelo pandas.

    Returns:
        dict[str, pd.DataFrame]:
            Dicionário com todas as planilhas importadas.

    Raises:
        ImportacaoDataFrameError:
            Caso alguma planilha não possa ser importada.
    """
    # Valida o arquivo e obtém os nomes disponíveis
    arquivo = validar_arquivo_excel(caminho)
    planilhas = listar_planilhas(arquivo)

    # Reutiliza a importação de múltiplas planilhas
    return importar_multiplas_planilhas(
        arquivo,
        planilhas=planilhas,
        linha_cabecalho=linha_cabecalho
    )


def importar_intervalo(
    caminho: str | Path,
    nome_planilha: str,
    intervalo: str,
    primeira_linha_cabecalho: bool = True
) -> pd.DataFrame:
    """
    Importa um intervalo específico do Excel como DataFrame.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str):
            Nome da planilha que contém o intervalo.
        intervalo (str):
            Referência do intervalo, como A1:D10.
        primeira_linha_cabecalho (bool):
            Define se a primeira linha do intervalo contém os cabeçalhos.

    Returns:
        pd.DataFrame:
            DataFrame contendo somente os dados do intervalo.

    Raises:
        IntervaloImportacaoInvalidoError:
            Caso o intervalo seja inválido.
        ImportacaoDataFrameError:
            Caso o arquivo não possa ser lido.
    """
    # Valida o arquivo, a planilha e o intervalo
    arquivo = validar_arquivo_excel(
        caminho,
        permitir_xls=False
    )
    planilha = _validar_planilha_existente(
        arquivo,
        nome_planilha
    )
    _, limites = _normalizar_intervalo(intervalo)

    # Valida a opção de cabeçalho
    _validar_booleano(
        primeira_linha_cabecalho,
        "primeira_linha_cabecalho"
    )

    # Extrai os limites do intervalo
    coluna_inicial, linha_inicial, coluna_final, linha_final = limites

    # Abre o Workbook em modo de leitura
    try:
        workbook = load_workbook(
            arquivo,
            read_only=True,
            data_only=False
        )
    except Exception as erro:
        raise ImportacaoDataFrameError(
            f"Não foi possível abrir o arquivo: {arquivo}"
        ) from erro

    try:
        # Obtém a planilha solicitada
        worksheet = workbook[planilha]

        # Cria a matriz com os valores do intervalo
        linhas = [
            [
                worksheet.cell(
                    row=linha,
                    column=coluna
                ).value
                for coluna in range(coluna_inicial, coluna_final + 1)
            ]
            for linha in range(linha_inicial, linha_final + 1)
        ]
    finally:
        # Fecha o Workbook após a leitura
        workbook.close()

    # Impede a criação de DataFrame sem linhas
    if not linhas:
        return pd.DataFrame()

    # Utiliza a primeira linha como cabeçalho quando solicitado
    if primeira_linha_cabecalho:
        colunas_dataframe = [
            str(valor) if valor is not None else f"Coluna_{indice}"
            for indice, valor in enumerate(linhas[0], start=1)
        ]
        dados = linhas[1:]
    else:
        colunas_dataframe = None
        dados = linhas

    # Cria e retorna o DataFrame do intervalo
    return pd.DataFrame(
        dados,
        columns=colunas_dataframe
    )


def importar_colunas(
    caminho: str | Path,
    nome_planilha: str,
    colunas: list[str],
    linha_cabecalho: int | None = LINHA_CABECALHO_PADRAO
) -> pd.DataFrame:
    """
    Importa colunas específicas de uma planilha Excel.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str):
            Nome da planilha importada.
        colunas (list[str]):
            Lista com nomes ou letras das colunas selecionadas.
        linha_cabecalho (int | None):
            Linha usada como cabeçalho pelo pandas.

    Returns:
        pd.DataFrame:
            DataFrame contendo somente as colunas selecionadas.

    Raises:
        TypeError:
            Caso colunas não seja uma lista de strings.
        ValueError:
            Caso a lista esteja vazia ou possua duplicidades.
    """
    # Verifica se as colunas foram informadas em uma lista
    if not isinstance(colunas, list):
        raise TypeError(
            "O parâmetro 'colunas' deve ser uma lista de strings."
        )

    # Verifica se a lista contém valores
    if not colunas:
        raise ValueError(
            "O parâmetro 'colunas' não pode estar vazio."
        )

    # Verifica se todos os itens são strings
    if not all(isinstance(coluna, str) for coluna in colunas):
        raise TypeError(
            "Todos os itens de 'colunas' devem ser strings."
        )

    # Verifica se existem itens duplicados
    if len(colunas) != len(set(colunas)):
        raise ValueError(
            "A lista 'colunas' não pode possuir valores duplicados."
        )

    # Importa apenas as colunas selecionadas
    return importar_dataframe(
        caminho,
        nome_planilha=nome_planilha,
        linha_cabecalho=linha_cabecalho,
        colunas=colunas
    )


def importar_registros(
    caminho: str | Path,
    nome_planilha: str | int = 0,
    linha_cabecalho: int | None = LINHA_CABECALHO_PADRAO,
    remover_valores_ausentes: bool = False
) -> list[dict[str, Any]]:
    """
    Importa uma planilha e retorna seus registros como dicionários.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str | int):
            Nome ou índice da planilha importada.
        linha_cabecalho (int | None):
            Linha usada como cabeçalho pelo pandas.
        remover_valores_ausentes (bool):
            Define se chaves com valores ausentes serão removidas.

    Returns:
        list[dict[str, Any]]:
            Lista contendo um dicionário para cada linha importada.

    Raises:
        TypeError:
            Caso algum parâmetro possua tipo incompatível.
        ImportacaoDataFrameError:
            Caso os dados não possam ser importados.
    """
    # Valida a opção de remoção dos valores ausentes
    _validar_booleano(
        remover_valores_ausentes,
        "remover_valores_ausentes"
    )

    # Importa a planilha como DataFrame
    dataframe = importar_dataframe(
        caminho,
        nome_planilha=nome_planilha,
        linha_cabecalho=linha_cabecalho
    )

    # Converte o DataFrame para registros
    registros = dataframe.to_dict(
        orient="records"
    )

    # Converte NaN e NaT para None
    registros_normalizados = [
        {
            chave: None if pd.isna(valor) else valor
            for chave, valor in registro.items()
        }
        for registro in registros
    ]

    # Remove chaves com valores ausentes quando solicitado
    if remover_valores_ausentes:
        registros_normalizados = [
            {
                chave: valor
                for chave, valor in registro.items()
                if valor is not None
            }
            for registro in registros_normalizados
        ]

    # Retorna os registros normalizados
    return registros_normalizados


# ----------------------------------------------------------------------------
# FUNÇÕES DE CABEÇALHOS E TABELAS
# ----------------------------------------------------------------------------

def importar_cabecalhos(
    caminho: str | Path,
    nome_planilha: str,
    linha_cabecalho: int = 1,
    ignorar_vazios: bool = True
) -> list[Any]:
    """
    Importa os valores de uma linha utilizada como cabeçalho.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str):
            Nome da planilha consultada.
        linha_cabecalho (int):
            Número da linha do Excel, iniciado em um.
        ignorar_vazios (bool):
            Define se cabeçalhos None serão removidos.

    Returns:
        list[Any]:
            Lista com os valores encontrados na linha.

    Raises:
        PlanilhaNaoEncontradaError:
            Caso a planilha não exista.
        ImportacaoDataFrameError:
            Caso o arquivo não possa ser lido.
    """
    # Valida o arquivo, a planilha, a linha e a opção
    arquivo = validar_arquivo_excel(
        caminho,
        permitir_xls=False
    )
    planilha = _validar_planilha_existente(
        arquivo,
        nome_planilha
    )
    linha = _validar_inteiro(
        linha_cabecalho,
        "linha_cabecalho",
        minimo=1
    )
    _validar_booleano(
        ignorar_vazios,
        "ignorar_vazios"
    )

    # Abre o Workbook em modo de leitura
    workbook = load_workbook(
        arquivo,
        read_only=True,
        data_only=False
    )

    try:
        # Obtém os valores da linha selecionada
        valores = [
            celula.value
            for celula in workbook[planilha][linha]
        ]
    finally:
        # Fecha o Workbook após a leitura
        workbook.close()

    # Remove valores None quando solicitado
    if ignorar_vazios:
        valores = [
            valor
            for valor in valores
            if valor is not None
        ]

    # Retorna os cabeçalhos encontrados
    return valores


def localizar_linha_cabecalho(
    caminho: str | Path,
    nome_planilha: str,
    cabecalhos_esperados: Iterable[str],
    limite_linhas: int = 20,
    considerar_maiusculas: bool = False
) -> int:
    """
    Localiza a primeira linha que contém os cabeçalhos esperados.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str):
            Nome da planilha consultada.
        cabecalhos_esperados (Iterable[str]):
            Coleção de cabeçalhos que devem existir na mesma linha.
        limite_linhas (int):
            Quantidade máxima de linhas analisadas.
        considerar_maiusculas (bool):
            Define se a comparação diferencia maiúsculas e minúsculas.

    Returns:
        int:
            Número da linha encontrada, iniciado em um.

    Raises:
        ValueError:
            Caso os cabeçalhos não sejam encontrados no limite informado.
    """
    # Rejeita strings isoladas
    if isinstance(cabecalhos_esperados, (str, bytes)) or not isinstance(
        cabecalhos_esperados,
        Iterable
    ):
        raise TypeError(
            "O parâmetro 'cabecalhos_esperados' deve ser um iterável não textual."
        )

    # Valida e normaliza os cabeçalhos esperados
    esperados = [
        _validar_texto(cabecalho, "cabecalho")
        for cabecalho in cabecalhos_esperados
    ]

    # Verifica se a coleção possui conteúdo
    if not esperados:
        raise ValueError(
            "O parâmetro 'cabecalhos_esperados' não pode estar vazio."
        )

    # Valida o limite e a opção de comparação
    limite = _validar_inteiro(
        limite_linhas,
        "limite_linhas",
        minimo=1
    )
    _validar_booleano(
        considerar_maiusculas,
        "considerar_maiusculas"
    )

    # Normaliza os valores esperados quando necessário
    conjunto_esperado = {
        valor if considerar_maiusculas else valor.casefold()
        for valor in esperados
    }

    # Analisa as linhas dentro do limite
    for numero_linha in range(1, limite + 1):
        # Obtém os valores preenchidos da linha
        valores = importar_cabecalhos(
            caminho,
            nome_planilha,
            linha_cabecalho=numero_linha,
            ignorar_vazios=True
        )

        # Normaliza os valores encontrados
        conjunto_encontrado = {
            str(valor) if considerar_maiusculas else str(valor).casefold()
            for valor in valores
        }

        # Retorna a primeira linha que contém todos os cabeçalhos
        if conjunto_esperado <= conjunto_encontrado:
            return numero_linha

    # Interrompe a execução quando nenhum cabeçalho foi localizado
    raise ValueError(
        "Os cabeçalhos esperados não foram encontrados nas primeiras "
        f"{limite} linhas."
    )


def importar_tabela_excel(
    caminho: str | Path,
    nome_planilha: str,
    nome_tabela: str
) -> pd.DataFrame:
    """
    Importa uma tabela estruturada do Excel como DataFrame.

    Args:
        caminho (str | Path):
            Caminho do arquivo Excel.
        nome_planilha (str):
            Nome da planilha que contém a tabela.
        nome_tabela (str):
            Nome interno da tabela estruturada.

    Returns:
        pd.DataFrame:
            DataFrame criado a partir do intervalo da tabela.

    Raises:
        TabelaExcelNaoEncontradaError:
            Caso a tabela não exista na planilha.
        ImportacaoDataFrameError:
            Caso o arquivo não possa ser lido.
    """
    # Valida o arquivo, a planilha e o nome da tabela
    arquivo = validar_arquivo_excel(
        caminho,
        permitir_xls=False
    )
    planilha = _validar_planilha_existente(
        arquivo,
        nome_planilha
    )
    tabela = _validar_texto(
        nome_tabela,
        "nome_tabela"
    )

    # Abre o Workbook em modo de leitura normal para consultar as tabelas
    try:
        workbook = load_workbook(
            arquivo,
            read_only=False,
            data_only=False
        )
    except Exception as erro:
        raise ImportacaoDataFrameError(
            f"Não foi possível abrir o arquivo: {arquivo}"
        ) from erro

    try:
        # Obtém a Worksheet correspondente
        worksheet = workbook[planilha]

        # Verifica se a tabela existe
        if tabela not in worksheet.tables:
            raise TabelaExcelNaoEncontradaError(
                f"A tabela '{tabela}' não foi encontrada na planilha "
                f"'{planilha}'."
            )

        # Obtém a referência da tabela estruturada
        referencia = worksheet.tables[tabela].ref
    finally:
        # Fecha o Workbook após a consulta
        workbook.close()

    # Importa o intervalo correspondente à tabela
    return importar_intervalo(
        arquivo,
        nome_planilha=planilha,
        intervalo=referencia,
        primeira_linha_cabecalho=True
    )


# ----------------------------------------------------------------------------
# EXPORTAÇÕES PÚBLICAS
# ----------------------------------------------------------------------------

__all__ = [
    "EXTENSOES_EXCEL",
    "EXTENSOES_OPENPYXL",
    "EXTENSAO_EXCEL_LEGADA",
    "ENGINE_EXCEL_PADRAO",
    "ENGINE_EXCEL_LEGADO",
    "LINHA_CABECALHO_PADRAO",
    "ExcelImportacaoError",
    "ArquivoExcelInvalidoError",
    "ArquivoExcelNaoEncontradoError",
    "PlanilhaNaoEncontradaError",
    "IntervaloImportacaoInvalidoError",
    "TabelaExcelNaoEncontradaError",
    "ImportacaoDataFrameError",
    "validar_arquivo_excel",
    "listar_planilhas",
    "planilha_existe",
    "obter_metadados_arquivo",
    "importar_dataframe",
    "importar_multiplas_planilhas",
    "importar_todas_planilhas",
    "importar_intervalo",
    "importar_colunas",
    "importar_registros",
    "importar_cabecalhos",
    "localizar_linha_cabecalho",
    "importar_tabela_excel"
]
